"""Bestandsimport (Phase 2 · Lieferung 3):  python tests/test_material_import.py  oder  pytest tests/test_material_import.py

Testplan 6.1 des Bauentwurfs L3 – angepasst an die Entscheidungen vom 19.10.2026: Zeilen ohne Seriennummer bekommen
ohne Haken eine STR-Nummer, offene Nacharbeit verfällt nach einem Monat; und vom 20.10.2026: Eine Liste ohne
Seriennummernspalte wird nur mit ausdrücklicher Bestätigung übernommen. Dateien entstehen im Test (CSV in mehreren
Kodierungen, xlsx mit openpyxl, ZIP mit zipfile). Alle Namen, Hersteller, Produkte und Nummern sind erfundene
Platzhalter („Testwerk“, „Gurt Alpha“, „OG Musterstadt“, „SN-0001“)."""
import csv
import io
import json
import os
import re
import sqlite3
import sys
import tempfile
import time
import zipfile
from datetime import date, datetime, timedelta, timezone
from types import SimpleNamespace

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from test_material import (H, _heute, freigeben, gliederung, katalog_grund, make_app, neuer_artikel,  # noqa: E402
                           neues_konto, umgebung)
from test_material_trennung import sql, zeilen, zwei_gliederungen  # noqa: E402
from test_material_ueberblick import ohne_post, ordner_neu  # noqa: E402

KOPF = ["Hersteller", "Modell", "Seriennummer"]


# --- Hilfen -------------------------------------------------------------------------------------------------------

def csv_bytes(zeilen_, kodierung="utf-8-sig", trenn=";"):
    buf = io.StringIO()
    csv.writer(buf, delimiter=trenn, lineterminator="\r\n").writerows(zeilen_)
    return buf.getvalue().encode(kodierung)


def xlsx_bytes(zeilen_, blaetter=None, formate=None):
    """Eine Arbeitsmappe; blaetter: weitere Blätter davor (Liste von Zeilenlisten); formate: {(zeile, spalte): Format}
    (1-basiert)."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for n, inhalt in enumerate((blaetter or []) + [zeilen_]):
        ws = wb.create_sheet(f"Blatt {n + 1}")
        for z in inhalt:
            ws.append(z)
    for (r, c), fmt in (formate or {}).items():
        ws.cell(row=r, column=c).number_format = fmt
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _senden(c, url, daten, name, felder):
    data = {"file": (io.BytesIO(daten), name)}
    for k, v in felder.items():
        data[k] = json.dumps(v) if isinstance(v, dict) else str(v)
    return c.post(url, headers=H, content_type="multipart/form-data", data=data)


def probelauf(c, fid, daten, name="liste.csv", **felder):
    return _senden(c, f"/api/material/ordner/{fid}/import/probelauf", daten, name, felder)


def uebernehmen(c, fid, daten, name="liste.csv", **felder):
    return _senden(c, f"/api/material/ordner/{fid}/import", daten, name, felder)


def ok(r, code=200):
    assert r.status_code == code, (r.status_code, r.get_json(silent=True))
    return r.get_json()


def nach_nr(bericht):
    return {z["nr"]: z for z in bericht["zeilen"]}


def codes(z):
    return [g["code"] for g in z["gruende"]]


def stand_aller_tabellen(app):
    """Zeilenzahlen aller mat_-Tabellen und mat_meta (samt Inventarzähler) – vor und nach einem Probelauf gleich."""
    with app.app_context():
        from app import db
        namen = [r["name"] for r in db.query("SELECT name FROM sqlite_master WHERE type = 'table' AND name LIKE "
                                              "'mat\\_%' ESCAPE '\\' AND name NOT LIKE '%fts%' ORDER BY name")]
        out = {t: db.query(f"SELECT COUNT(*) AS n FROM {t}", one=True)["n"] for t in namen}
        out["mat_meta"] = db.query("SELECT * FROM mat_meta ORDER BY schluessel")
        return out


def alle_texte(app):
    """Alles, was in den mat_-Tabellen steht, als ein Text (für „steht nirgends“)."""
    with app.app_context():
        from app import db
        namen = [r["name"] for r in db.query("SELECT name FROM sqlite_master WHERE type = 'table' AND name LIKE "
                                              "'mat\\_%' ESCAPE '\\' ORDER BY name")]
        return json.dumps({t: db.query(f"SELECT * FROM {t}") for t in namen}, ensure_ascii=False, default=str)


def tag(d):
    return d.strftime("%d.%m.%Y")


def plus(d, monate):
    from app.material.services.interval import plus_monate
    return plus_monate(d, monate)


def aufbau(tmp):
    """Umgebung, Katalog (Gurt Alpha textil, HMS Beta Metall), Gliederung mit Lager."""
    app, k = umgebung(tmp)
    ids = katalog_grund(k["adm"])
    gid, system = gliederung(k["adm"])
    lager = ordner_neu(k["wart"], gid, "Lager Musterstadt")
    return SimpleNamespace(app=app, k=k, ids=ids, gid=gid, system=system, lager=lager, w=k["wart"], adm=k["adm"],
                           mv=k["katalog"])


def produkt(mv, kategorie, hersteller, modell, nr, **felder):
    d = dict({"category_id": kategorie, "manufacturer": hersteller, "model": modell, "article_no": nr,
              "material_kind": "metall_hardware", "interval_months": 12}, **felder)
    return ok(mv.post("/api/material/produkte", json=d), 201)["produkt"]["id"]


def offene_zeilen(app, fid):
    return zeilen(app, "SELECT r.* FROM mat_import_row r JOIN mat_import i ON i.id = r.import_id WHERE i.folder_id = ? "
                       "AND r.status = 'offen' ORDER BY r.id", (fid,))


def vereinsliste(p, T):
    """Eine typische Vereinsliste: Titelzeilen, Personenspalte, Bezeichnung links vom Typ, Größe und Standort,
    Baujahr nur als Jahr (Platzhalterwerte). p: Vorsatz der Seriennummern."""
    letzte = tag(plus(T, -7))
    naechste = tag(plus(T, 3))
    return [["Bestandsliste OG Musterstadt (Platzhalter)"], ["Stand: 01.09.2026"], [],
            ["Lfd. Nr.", "Bezeichnung", "Hersteller", "Typ", "Größe", "Seriennummer", "Baujahr", "Kaufdatum",
             "letzte Prüfung", "nächste Prüfung", "Standort", "Ausgegeben an", "Anzahl", "Bemerkung"],
            [1, "Karabiner", "Testwerk", "HMS Beta", "", f"{p}-0001", "2019", "05/2019", letzte, "", "MTW",
             "Platzhalter Person A", 1, "Platzhalter-Bemerkung"],
            [2, "Gurt", "Testwerk", "Gurt Alpha", "M", f"{p}-0002", "03/2019", "", letzte, "", "MTW", "", 1, ""],
            [3, "Gurt", "Testwerk", "Gurt Alpha", "L", f"{p}-0003", "2019", "", "", "", "Lager", "", 1, ""],
            [4, "Karabiner", "Testwerk", "HMS Beta", "", f"{p}-0004", "", "", "", naechste, "Lager", "", "", ""],
            [5, "Helm", "Kopfwerk", "Platzhalterhelm", "", f"{p}-0005", "", "", "", "", "MTW", "", 1, ""],
            [6, "Karabiner", "Testwerk", "HMS Beta", "", f"{p}-0006", "", "", "", "", "MTW", "Platzhalter Person B", 1,
             "Größe egal"],
            [7, "Karabiner", "Testwerk", "HMS Beta", "", "", "", "", "", "", "", "", 1, ""]]


# --- 1. Migration -------------------------------------------------------------------------------------------------

def test_migration_12():
    from app.material import models
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        neuer_artikel(t.w, t.ids["hms"], t.lager, serial_no="MIG-0001")
        ok(t.mv.post("/api/material/formulare", json={"kategorie_id": t.ids["gurte"], "titel": "Platzhalterformular",
                                                       "punkte": [{"text": "Testpunkt", "art": "yesno", "soll": "ja"}]}),
           201)
        pfad = os.path.join(tmp, "stroemis.db")
        conn = sqlite3.connect(pfad)
        assert conn.execute("SELECT wert FROM mat_meta WHERE schluessel = 'schema_version'").fetchone()[0] == "12"
        spalten = {t_: [r[1] for r in conn.execute(f"PRAGMA table_info({t_})")] for t_ in ("mat_import", "mat_import_row")}
        assert spalten["mat_import"] == ["id", "folder_id", "file_name", "file_sha256", "file_kind", "rows_total",
                                         "rows_created", "rows_existing", "rows_rework", "created_by", "created_by_name",
                                         "created_at"], spalten
        assert spalten["mat_import_row"] == ["id", "import_id", "line_no", "values_json", "reason_code", "reason",
                                             "offer_key", "product_id", "request_id", "status", "item_id", "resolved_by",
                                             "resolved_by_name", "resolved_at"], spalten
        assert "wiki_page_id" in [r[1] for r in conn.execute("PRAGMA table_info(mat_form)")]
        indizes = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type = 'index'")}
        assert {"idx_mat_import_folder", "idx_mat_import_datei", "idx_mat_import_row_status", "idx_mat_import_row_angabe",
                "idx_mat_import_row_meldung", "idx_mat_form_wiki"} <= indizes
        for t_ in ("mat_import", "mat_import_row"):
            assert "AUTOINCREMENT" in conn.execute("SELECT sql FROM sqlite_master WHERE name = ?", (t_,)).fetchone()[0]
        # CHECKs: Dateiart, Status (mit „ersetzt“), offen ⇔ ohne resolved_at, line_no ≥ 1, UNIQUE(import_id, line_no).
        conn.execute("PRAGMA foreign_keys = ON")
        conn.execute("INSERT INTO mat_import (folder_id, file_kind, created_at) VALUES (?, 'csv', 'x')", (t.lager,))
        imp = conn.execute("SELECT MAX(id) FROM mat_import").fetchone()[0]
        for befehl, args in (("INSERT INTO mat_import (folder_id, file_kind, created_at) VALUES (?, 'xls', 'x')",
                              (t.lager,)),
                             ("INSERT INTO mat_import_row (import_id, line_no, status) VALUES (?, 0, 'offen')", (imp,)),
                             ("INSERT INTO mat_import_row (import_id, line_no, status) VALUES (?, 2, 'unbekannt')", (imp,)),
                             ("INSERT INTO mat_import_row (import_id, line_no, status) VALUES (?, 2, 'verworfen')", (imp,)),
                             ("INSERT INTO mat_import_row (import_id, line_no, status, resolved_at) VALUES "
                              "(?, 2, 'offen', 'x')", (imp,))):
            try:
                conn.execute(befehl, args)
                raise AssertionError(befehl)
            except sqlite3.IntegrityError:
                pass
        conn.execute("INSERT INTO mat_import_row (import_id, line_no, status, resolved_at) VALUES (?, 2, 'ersetzt', 'x')",
                     (imp,))
        try:
            conn.execute("INSERT INTO mat_import_row (import_id, line_no) VALUES (?, 2)", (imp,))
            raise AssertionError("UNIQUE")
        except sqlite3.IntegrityError:
            pass
        conn.rollback()

        # 11 → 12: zuerst der Teilindex, dann die Spalte, dann die Tabellen.
        def stand():
            c = sqlite3.connect(pfad)
            out = {"form": c.execute("SELECT id, category_id, product_id, version, title, valid_from, active, created_at, "
                                     "updated_at FROM mat_form ORDER BY id").fetchall(),
                   "item": c.execute("SELECT * FROM mat_item ORDER BY id").fetchall(),
                   "product": c.execute("SELECT * FROM mat_product ORDER BY id").fetchall()}
            c.close()
            return out
        vorher = stand()
        for befehl in ("DROP INDEX idx_mat_form_wiki", "ALTER TABLE mat_form DROP COLUMN wiki_page_id",
                       "DROP TABLE mat_import_row", "DROP TABLE mat_import",
                       "UPDATE mat_meta SET wert = '11' WHERE schluessel = 'schema_version'"):
            conn.execute(befehl)
        conn.commit()
        conn.close()
        make_app(tmp)
        conn = sqlite3.connect(pfad)
        assert conn.execute("SELECT wert FROM mat_meta WHERE schluessel = 'schema_version'").fetchone()[0] == "12"
        assert {"mat_import", "mat_import_row"} <= {r[0] for r in conn.execute("SELECT name FROM sqlite_master")}
        assert conn.execute("SELECT COUNT(*) FROM mat_form WHERE wiki_page_id IS NOT NULL").fetchone()[0] == 0
        assert stand() == vorher
        # Zweiter Lauf: keine Änderung (auch Schritt 12 selbst nicht).
        assert models.migrieren(conn) == 12
        conn.execute("BEGIN")
        n = conn.total_changes
        models._schritt_12(conn)
        assert conn.total_changes == n
        conn.rollback()
        # Kaskade und SET NULL: Produkt, Meldung, Artikel weg → NULL; Ordner weg → Import und Zeilen weg.
        conn.execute("PRAGMA foreign_keys = ON")
        ordner = conn.execute("INSERT INTO mat_folder (parent_id, name, created_at, updated_at) VALUES (?, 'Weg', 'x', "
                              "'x')", (t.gid,)).lastrowid
        pid = conn.execute("INSERT INTO mat_product (category_id, manufacturer, model, dedupe_key, material_kind, "
                           "interval_months, created_at, updated_at) VALUES (?, 'Wegwerk', 'Weg', 'wegwerk|weg|', "
                           "'metall_hardware', 12, 'x', 'x')", (t.ids["karabiner"],)).lastrowid
        mid = conn.execute("INSERT INTO mat_product_request (manufacturer, model, dedupe_key, created_at) VALUES "
                           "('Wegwerk', 'Weg', 'x', 'x')").lastrowid
        imp = conn.execute("INSERT INTO mat_import (folder_id, file_kind, created_at) VALUES (?, 'csv', 'x')",
                           (ordner,)).lastrowid
        conn.execute("INSERT INTO mat_import_row (import_id, line_no, product_id, request_id) VALUES (?, 2, ?, ?)",
                     (imp, pid, mid))
        conn.execute("DELETE FROM mat_product WHERE id = ?", (pid,))
        conn.execute("DELETE FROM mat_product_request WHERE id = ?", (mid,))
        assert conn.execute("SELECT product_id, request_id FROM mat_import_row WHERE import_id = ?", (imp,)).fetchone() \
            == (None, None)
        conn.execute("DELETE FROM mat_folder WHERE id = ?", (ordner,))
        assert not conn.execute("SELECT 1 FROM mat_import WHERE id = ?", (imp,)).fetchone()
        assert not conn.execute("SELECT 1 FROM mat_import_row WHERE import_id = ?", (imp,)).fetchone()
        conn.rollback()
        conn.close()
        # Die genutzten Namen aus pruefungen_import gibt es (ein Umbau dort fiele hier auf).
        import app.pruefungen_import as pi
        from app.material.services import bestandsimport
        assert pi._arbeitsblatt(xlsx_bytes([["Hersteller", "Modell"], ["Testwerk", "HMS Beta"]]), "") == [
            ["Hersteller", "Modell"], ["Testwerk", "HMS Beta"]]
        assert issubclass(pi.ImportFehler, ValueError) and pi.normalisiert(" Hersteller\n(Firma) ") == "hersteller (firma)"
        assert pi._beginnt_mit("hersteller-nr.", "hersteller-nr") and pi.datum_lesen("15.03.21") == "2021-03-15"
        assert pi.als_text(3.0) == "3" and pi.als_text(datetime(2021, 3, 1)) == "2021-03-01"
        assert pi.MAX_DATENZEILEN == bestandsimport.MAX_DATENZEILEN == 2000 and pi.MAX_SPALTEN == 300
        assert pi.MAX_KOPF_SUCHE == 40 and pi.MAX_ZIP_EINTRAEGE == 500
    print("Migration Schritt 12 bestanden.")


# --- 2. Vereinsliste ----------------------------------------------------------------------------------------------

def test_vereinsliste_ende_zu_ende():
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        T = _heute()
        for art, daten, name in (("xlsx", xlsx_bytes(vereinsliste("VX", T)), "Bestand.xlsx"),
                                 ("csv", csv_bytes(vereinsliste("VC", T), "cp1252"), "Bestand.csv")):
            b = ok(probelauf(t.w, t.lager, daten, name))
            assert b["datei"]["art"] == art and b["datei"]["kopfzeile"] == 4 and b["datei"]["datenzeilen"] == 7, b["datei"]
            if art == "csv":
                assert (b["datei"]["kodierung"], b["datei"]["trennzeichen"]) == ("cp1252", ";")
            zu = {s["titel"]: (s["zuordnung"], s["gesperrt"]) for s in b["spalten"]}
            assert zu == {"Lfd. Nr.": ("ignorieren", False), "Bezeichnung": ("ignorieren", False),
                          "Hersteller": ("hersteller", False), "Typ": ("modell", False), "Größe": ("bemerkung", False),
                          "Seriennummer": ("seriennummer", False), "Baujahr": ("hergestellt", False),
                          "Kaufdatum": ("gekauft", False), "letzte Prüfung": ("zuletzt_geprueft", False),
                          "nächste Prüfung": ("naechste_pruefung", False), "Standort": ("bemerkung", False),
                          "Ausgegeben an": ("ignorieren", True), "Anzahl": ("anzahl", False),
                          "Bemerkung": ("bemerkung", False)}, zu
            assert not b["fehlt"] and any("„Bezeichnung“ wird nicht als Modell verwendet" in w_ for w_ in b["warnungen"])
            assert [s["beispiele"] for s in b["spalten"] if s["gesperrt"]] == [[]]
            z = nach_nr(b)
            assert b["zahlen"]["uebernehmbar"] == 5 and b["zahlen"]["nacharbeit"] == 2, b["zahlen"]
            assert codes(z[7]) == ["datum"] and codes(z[9]) and codes(z[9])[0].startswith("produkt_"), (z[7], z[9])
            assert b["zahlen"]["ohne_serie"] == 1 and any("STR-Nummer" in w_ for w_ in b["warnungen"])
            # Einzahl, wenn es nur eine Zeile ist (Durchsicht der Oberfläche vom 19.10.2026: „1 Zeilen“).
            assert ("1 Zeile ohne Seriennummer – bekommt eine STR-Nummer vom System (bitte am Artikel anbringen)."
                    in b["warnungen"]), b["warnungen"]
            r = ok(uebernehmen(t.w, t.lager, daten, name), 201)
            assert r["import"]["angelegt"] == 5 and r["import"]["nacharbeit"] == 2, r["import"]
            assert any(w_.startswith("1 Artikel ohne Herstellerseriennummer hat eine Seriennummer vom System")
                       for w_ in r["warnungen"]), r["warnungen"]
            z = nach_nr(r)
            with t.app.app_context():
                from app import db
                notes = {n: db.query("SELECT notes FROM mat_item WHERE id = ?", (z[n]["artikel"]["id"],),
                                     one=True)["notes"] for n in (5, 6, 10)}
            assert notes[5] == "Standort: MTW\nPlatzhalter-Bemerkung\nHerstelldatum laut Liste: 2019", notes
            assert notes[6] == "Größe: M\nStandort: MTW", notes
            assert notes[10] == "Standort: MTW\nGröße egal", notes
            assert z[11]["seriennummer"].startswith("STR-") and z[11]["artikel"]["inventory_no"].startswith("M-")
        texte = alle_texte(t.app)
        assert "Platzhalter Person" not in texte and "Platzhalter-Bemerkung" in texte
    print("Vereinsliste Ende zu Ende bestanden.")


# --- 3./4. Lesen ----------------------------------------------------------------------------------------------------

def _abgelehnt(f, *a, code=400):
    from app.material.services import Abgelehnt
    try:
        f(*a)
    except Abgelehnt as exc:
        assert exc.code == code, (exc.code, str(exc))
        return str(exc)
    raise AssertionError(f"keine Ablehnung: {a!r}"[:200])


def test_lesen_csv():
    from app.material.services import bestandsimport, tabelle
    zeilen_ = [["Hersteller", "Modell", "Ordner"], ["Prüfwerk Ä", "Gurt Ö", "Fächer Ü"]]
    for kodierung, trenn in (("utf-8-sig", ";"), ("utf-8", ","), ("cp1252", ";"), ("utf-16", "\t")):
        tab = tabelle.lesen(csv_bytes(zeilen_, kodierung, trenn), "liste.csv")
        assert (tab.art, tab.kodierung, tab.trennzeichen) == ("csv", kodierung, trenn), (tab.kodierung, kodierung)
        assert tab.zeilen == zeilen_, tab.zeilen
    # Anführungszeichen mit Trennzeichen und Zeilenumbruch im Feld; Titelzeilen; leere Zeilen.
    roh = ('Bestandsliste\r\n\r\nHersteller;Modell;Bemerkung\r\nTestwerk;"HMS; Beta";"Zeile eins\r\nZeile zwei"\r\n'
           '\r\n;;\r\nTestwerk;Gurt Alpha;\r\n').encode("utf-8")
    tab = tabelle.lesen(roh, "liste.csv")
    sp, nacharbeit_csv, kopf, _ = bestandsimport.spalten(tab)
    assert kopf == 2 and not nacharbeit_csv and tab.zeilen[3] == ["Testwerk", "HMS; Beta", "Zeile eins\r\nZeile zwei"]
    daten, leer, ohne = bestandsimport.datensaetze(tab, sp, kopf)
    assert [d["werte"] for d in daten] == [{"hersteller": "Testwerk", "modell": "HMS; Beta",
                                            "bemerkung": "Zeile eins\nZeile zwei"},
                                           {"hersteller": "Testwerk", "modell": "Gurt Alpha"}] and leer == 2, daten
    assert "Steuerzeichen" in _abgelehnt(tabelle.lesen, b"Hersteller;Modell\r\nA\x00B;C\r\n", "x.csv")
    assert "Zeichensatz nicht erkannt" in _abgelehnt(tabelle.lesen, "Hersteller;Modell\r\nWärk;X\x81\r\n".encode(
        "latin-1"), "x.csv")
    assert "Zeichensatz nicht erkannt" in _abgelehnt(tabelle.lesen, b"\xff\xfeH\x00e\x00x", "x.csv")
    assert "alte Excel-Datei" in _abgelehnt(tabelle.lesen, b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100, "x.csv")
    assert "PDF" in _abgelehnt(tabelle.lesen, b"%PDF-1.4 ...", "liste.csv")
    assert "Bild" in _abgelehnt(tabelle.lesen, b"\x89PNG\r\n\x1a\n" + b"\x00" * 20, "liste.csv")
    assert "leer" in _abgelehnt(tabelle.lesen, b"", "liste.csv")
    for name in ("liste.xls", "liste.ods"):
        assert "Excel-Arbeitsmappe (.xlsx)" in _abgelehnt(tabelle.lesen, b"a;b\r\n", name)
    # Mehr als 10 000 physische Zeilen: Abbruch, bevor alles gelesen ist – in Zeile 10 005 stünde ein Feld über der
    # Grenze des csv-Moduls, das sonst „unlesbar (Zeile 10005)“ ergäbe.
    roh = b"Hersteller;Modell\r\n" + b"Testwerk;HMS\r\n" * 10003 + b'"' + b"x" * 140000 + b'";y\r\n'
    assert "mehr als 10.000 Zeilen" in _abgelehnt(tabelle.lesen, roh, "x.csv")
    assert "unlesbar (Zeile 4)" in _abgelehnt(tabelle.lesen, b"Hersteller;Modell\r\nA;B\r\nC;D\r\n\"" + b"x" * 140000
                                              + b'";y\r\n', "x.csv")
    tab = tabelle.lesen(b"Hersteller;Modell\r\n" + b"Testwerk;HMS\r\n" * 2001, "x.csv")
    assert "höchstens 2000" in _abgelehnt(bestandsimport.spalten, tab)
    tab = tabelle.lesen(b"Hersteller;Modell\r\n" + b"Testwerk;HMS\r\n" * 2000 + b"\r\n" * 50, "x.csv")
    assert bestandsimport.spalten(tab)[2] == 0
    assert "höchstens 300" in _abgelehnt(tabelle.lesen, b"Hersteller;Modell\r\n" + b";".join([b"x"] * 301) + b"\r\n",
                                         "x.csv")
    # Formelschutz des eigenen Exports zurück; ein ' vor anderem Text bleibt.
    assert tabelle.zelltext("'=A1") == "=A1" and tabelle.zelltext("'-5") == "-5" and tabelle.zelltext("'abc") == "'abc"
    assert tabelle.zelltext("a\xa0 b") == "a b" and tabelle.zelltext(3.0) == "3"
    print("CSV lesen bestanden.")


def test_lesen_xlsx_und_grenzen():
    import app.pruefungen_import as pi
    from app.material.services import tabelle
    # Datumszelle und Zelle im Monatsformat (Excel speichert den vollen Tag).
    roh = xlsx_bytes([["Hersteller", "Herstelldatum", "Lebensende"], ["Testwerk", datetime(2021, 3, 1),
                                                                     datetime(2031, 3, 15)]],
                     formate={(2, 2): "mm/yyyy"})
    tab = tabelle.lesen(roh, "liste.xlsx")
    assert tab.art == "xlsx" and [tabelle.zelltext(w) for w in tab.zeilen[1]] == ["Testwerk", "2021-03-01", "2031-03-15"]
    # Erstes Blatt leer, zweites mit Daten.
    tab = tabelle.lesen(xlsx_bytes([["Hersteller", "Modell"], ["Testwerk", "HMS Beta"]], blaetter=[[]]), "liste.xlsx")
    assert tab.zeilen == [["Hersteller", "Modell"], ["Testwerk", "HMS Beta"]], tab.zeilen
    # ZIP mit 11 MB entpackt: abgelehnt, bevor pruefungen_import etwas öffnet (Spion).
    aufgerufen = []
    alt = pi._arbeitsblatt
    pi._arbeitsblatt = lambda *a: aufgerufen.append(a) or alt(*a)
    try:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("xl/worksheets/sheet1.xml", b"\x00" * (11 * 1024 * 1024))
        assert len(buf.getvalue()) < 100_000
        text = _abgelehnt(tabelle.lesen, buf.getvalue(), "liste.xlsx")
        assert "entpackt 11 MB" in text and "höchstens 10 MB" in text, text
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            for n in range(501):
                zf.writestr(f"x{n}.xml", b"")
        text = _abgelehnt(tabelle.lesen, buf.getvalue(), "liste.xlsx")
        assert "Bestandsliste" in text and "Teilnehmerliste" not in text, text
        assert not aufgerufen
        # Eine Meldung aus pruefungen_import nennt die Bestandsliste (21 Blätter).
        text = _abgelehnt(tabelle.lesen, xlsx_bytes([["a"]], blaetter=[[["x"]]] * 20), "liste.xlsx")
        assert "Bestandsliste" in text and "Teilnehmerliste" not in text and aufgerufen, text
    finally:
        pi._arbeitsblatt = alt
    assert "nicht als Excel-Arbeitsmappe" in _abgelehnt(tabelle.lesen, b"PK\x03\x04kaputt" + b"\x00" * 50, "l.xlsx")
    # Verbundene Zellen, Links, Kommentare über große Bereiche und dichte leere Zellen (Prüfung vor Auslieferung,
    # Sicherheit 1 – vorher bis 2,8 GB und 76 s): abgelehnt, bevor openpyxl lädt (pi._blaetter_vorpruefen), auch
    # verkleidet; über die Route in unter 2 s.
    from test_pruefungen_import import xlsx_bomben
    geladen = []
    alt_laden = pi.openpyxl.load_workbook
    pi.openpyxl.load_workbook = lambda *a, **k: geladen.append(1) or alt_laden(*a, **k)
    try:
        bomben = xlsx_bomben()
        for name, roh in bomben.items():
            text = _abgelehnt(tabelle.lesen, roh, f"{name}.xlsx")
            assert not geladen and "Teilnehmerliste" not in text, (name, text)
        with tempfile.TemporaryDirectory() as tmp:
            t = aufbau(tmp)
            for name in ("merge_6m", "link_6m", "merge_riesig", "leere_c"):
                t0 = time.time()
                r = probelauf(t.w, t.lager, bomben[name], f"{name}.xlsx")
                assert r.status_code == 400 and time.time() - t0 < 2 and not geladen, (name, r.get_json())
        tab = tabelle.lesen(xlsx_bytes([["Bestandsliste (Platzhalter)"], KOPF, ["Testwerk", "HMS Beta", "SN-0001"]]),
                            "liste.xlsx")
        assert tab.zeilen[2] == ["Testwerk", "HMS Beta", "SN-0001"] and geladen
    finally:
        pi.openpyxl.load_workbook = alt_laden
    # 5 MB + 1 Byte: 413 (über dem Dienst), deutlich mehr: 413 schon von der Upload-Grenze der Route.
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        r = probelauf(t.w, t.lager, b"Hersteller;Modell\r\n" + b"x" * (5 * 1024 * 1024), "gross.csv")
        assert r.status_code == 413 and "größer als 5 MB" in r.json["error"], r.get_json()
        r = probelauf(t.w, t.lager, b"x" * (7 * 1024 * 1024), "riesig.csv")
        assert r.status_code == 413 and "Grenze: 6 MB" in r.json["error"], r.get_json()
    print("xlsx lesen und Grenzen bestanden.")


# --- 5. Seriennummer ------------------------------------------------------------------------------------------------

def test_seriennummer():
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        werte = [123456789012, 1.2345678901234567e19, datetime(2026, 1, 2), "1,23457E+11", "1.2E+5", "19E0123456",
                 "000123", "01.02.2026", "STR-000012", "X" * 81, "17", "M-000017"]
        daten = xlsx_bytes([KOPF] + [["Testwerk", "HMS Beta", w] for w in werte])
        r = probelauf(t.w, t.lager, daten, "serien.xlsx")
        b = ok(r)
        assert "12345678901234567168" not in r.get_data(as_text=True)
        z = nach_nr(b)
        assert z[2]["status"] == "uebernehmbar" and z[2]["seriennummer"] == "123456789012"
        assert any("führende Nullen prüfen" in h for h in z[2]["hinweise"]), z[2]
        for n in (3, 4, 5, 6, 9, 10, 11):
            assert codes(z[n]) == ["serie_unbrauchbar"], (n, z[n])
        assert "15 Stellen" in z[3]["gruende"][0]["text"] and "Datum" in z[4]["gruende"][0]["text"]
        assert z[7]["status"] == "uebernehmbar" and z[8]["seriennummer"] == "000123" and z[8]["status"] == "uebernehmbar"
        assert codes(z[12]) == ["serie_von_hand"] and codes(z[13]) == ["serie_von_hand"], (z[12], z[13])
    print("Seriennummern bestanden.")


# --- 6. Spalten -----------------------------------------------------------------------------------------------------

def test_spalten():
    from app.material.services import bestandsimport, tabelle
    erkannt = {"Hersteller": "hersteller", "Marke": "hersteller", "Modell": "modell", "Typ": "modell",
               "Bezeichnung": "modell", "Artikel-Nr.": "artikelnummer", "Art.-Nr.": "artikelnummer",
               "Seriennummer": "seriennummer", "S/N": "seriennummer", "Hersteller-Seriennummer": "seriennummer",
               "Fabrik-Nr.": "seriennummer", "Inventar-Nr.": "inventarnummer", "Unterordner": "ordner",
               "Herstelldatum": "hergestellt", "Baujahr": "hergestellt", "Anschaffung": "gekauft",
               "Inbetriebnahme": "erstgebrauch", "Prüfdatum": "zuletzt_geprueft", "Letzte Pruefung": "zuletzt_geprueft",
               "Fällig am": "naechste_pruefung", "Prüfintervall": "intervall", "Ablaufdatum": "lebensende",
               "Stückzahl": "anzahl", "Notiz": "bemerkung", "Größe": "bemerkung", "Standort": "bemerkung",
               "Hersteller-Nr.": "ignorieren", "Lfd. Nr.": "ignorieren", "Prüffrist": "ignorieren",
               "Bestellnummer": "ignorieren", "Name": "ignorieren", "Kategorie": "ignorieren"}
    for titel, feld in erkannt.items():
        assert bestandsimport.erkennen(titel)[0] == feld, (titel, bestandsimport.erkennen(titel))
    for titel in ("Ausgegeben an", "Telefon", "E-Mail", "Vorname", "Träger"):
        assert bestandsimport.erkennen(titel) == ("ignorieren", 0, True), titel

    def lauf(kopf, zeile, zuordnung=None):
        tab = tabelle.Tabelle("csv", "utf-8", ";", [kopf, zeile], "x")
        sp, _, k, warn = bestandsimport.spalten(tab, zuordnung)
        daten, _, _ = bestandsimport.datensaetze(tab, sp, k)
        return {s["titel"]: s for s in sp}, warn, daten, bestandsimport.fehlt(sp)
    sp, warn, _, _ = lauf(["Bezeichnung", "Hersteller", "Typ"], ["Helm", "Testwerk", "HMS Beta"])
    assert sp["Typ"]["zuordnung"] == "modell" and sp["Bezeichnung"]["zuordnung"] == "ignorieren"
    assert any("„Bezeichnung“ wird nicht als Modell verwendet – „Typ“ hat Vorrang" in w for w in warn), warn
    sp, warn, _, _ = lauf(["Typ", "Hersteller", "Modell"], ["A", "Testwerk", "HMS Beta"])
    assert sp["Modell"]["zuordnung"] == "modell" and sp["Typ"]["zuordnung"] == "ignorieren"
    # Zwei Bemerkungsspalten und „Titel: Wert“; eine Personenspalte bleibt gesperrt, auch über die Zuordnung.
    sp, warn, daten, fehlt = lauf(["Hersteller", "Modell", "Größe", "Bemerkung", "Telefon", "Farbe"],
                                  ["Testwerk", "HMS Beta", "M", "neu", "0000 000", "rot"], {"4": "bemerkung"})
    assert daten[0]["werte"]["bemerkung"] == "Größe: M\nneu\nFarbe: rot", daten
    assert sp["Telefon"]["gesperrt"] and sp["Telefon"]["zuordnung"] == "ignorieren" and sp["Telefon"]["beispiele"] == []
    assert any("„Telefon“ enthält Personenangaben" in w for w in warn) and not fehlt
    # Auch in der Nacharbeits-CSV (feste Zuordnung): eine ergänzte Personenspalte bleibt gesperrt, auch über die Zuordnung.
    sp, warn, daten, _ = lauf(["Nacharbeit-Nr.", "Hersteller", "Modell", "Seriennummer", "Ausgegeben an"],
                              ["", "Testwerk", "HMS Beta", "A-1", "Platzhalter Person"], {"4": "bemerkung"})
    assert sp["Ausgegeben an"]["gesperrt"] and sp["Ausgegeben an"]["zuordnung"] == "ignorieren"
    assert sp["Ausgegeben an"]["beispiele"] == [] and "bemerkung" not in daten[0]["werte"], daten
    assert any("„Ausgegeben an“ enthält Personenangaben" in w for w in warn), warn
    # Ein Feld zweimal (außer Bemerkung): die erste gewinnt; die Zuordnung der Anfrage überschreibt.
    sp, warn, daten, _ = lauf(["Hersteller", "Modell", "Seriennummer", "S/N"], ["Testwerk", "HMS", "A-1", "B-2"])
    assert sp["S/N"]["zuordnung"] == "ignorieren" and daten[0]["werte"]["seriennummer"] == "A-1"
    assert any("„S/N“ ist ein zweites Mal als „Seriennummer“ zugeordnet" in w for w in warn), warn
    sp, warn, daten, _ = lauf(["Hersteller", "Modell", "Seriennummer", "S/N"], ["Testwerk", "HMS", "A-1", "B-2"],
                              {"3": "seriennummer", "2": "inventarnummer", "9": "modell", "x": "modell", "0": "unsinn"})
    assert daten[0]["werte"]["seriennummer"] == "B-2" and daten[0]["werte"]["inventarnummer"] == "A-1"
    assert len([w for w in warn if "ignoriert" in w]) == 3, warn
    _, _, _, fehlt = lauf(["Modell", "Seriennummer"], ["HMS", "A-1"])
    assert fehlt == ["hersteller"]
    _, _, _, fehlt = lauf(["Hersteller", "Seriennummer"], ["Testwerk", "A-1"])
    assert fehlt == ["modell"]
    print("Spalten bestanden.")


# --- 7. Daten und was blockiert -------------------------------------------------------------------------------------

def test_daten_und_blockieren():
    from app.material.services import tabelle
    for text, art, soll in (("15.03.2021", "monat", ("2021-03-15", None, None)), ("15.3.21", "tag", ("2021-03-15", None, None)),
                            ("2021-03-15 10:30", "tag", ("2021-03-15", None, None)),
                            ("03/2021", "monat", ("2021-03", None, None)), ("3/2021", "monat", ("2021-03", None, None)),
                            ("03.2021", "monat", ("2021-03", None, None)), ("2021-03", "monat", ("2021-03", None, None)),
                            ("03-2021", "tag", ("2021-03-01", None, "Monatsangabe – als 1. des Monats gelesen")),
                            ("2021", "monat", (None, "nur Jahr", None)), ("2021", "tag", (None, "nur Jahr", None)),
                            ("März 2021", "monat", (None, "unlesbar", None)), ("31.02.2021", "tag", (None, "unlesbar", None)),
                            ("03/21", "monat", (None, "unlesbar", None)), ("44270", "tag", (None, "unlesbar", None)),
                            ("13/2021", "monat", (None, "unlesbar", None)), ("", "tag", (None, None, None))):
        assert tabelle.datum(text, art) == soll, (text, art, tabelle.datum(text, art))
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        T = _heute()
        kopf = ["Hersteller", "Modell", "Seriennummer", "Herstelldatum", "Kaufdatum", "Erstgebrauch", "Zuletzt geprüft",
                "Prüfintervall", "Lebensende"]
        monat = f"{T.month:02d}/{T.year}"
        b = ok(probelauf(t.w, t.lager, csv_bytes([kopf,
            ["Testwerk", "Gurt Alpha", "DB-0001", "Baujahr 2019", "", "", "", "", ""],
            ["Testwerk", "HMS Beta", "DB-0002", "Baujahr 2019", "", "", "", "", ""],
            ["Testwerk", "HMS Beta", "DB-0003", "", "Frühjahr 2018", "", "", "", ""],
            ["Testwerk", "HMS Beta", "DB-0004", "", "", "", "Frühjahr 23", "", ""],
            ["Testwerk", "HMS Beta", "DB-0005", "", "", "", "", "ca. jährlich", ""],
            ["Testwerk", "Gurt Alpha", "DB-0006", "03/2020", "", "", "", "", "irgendwann"],
            ["Testwerk", "HMS Beta", "DB-0007", "", "", "", "", "", "irgendwann"],
            ["Testwerk", "HMS Beta", "DB-0008", "", "", tag(T + timedelta(days=40)), "", "", ""],
            ["Testwerk", "HMS Beta", "DB-0009", monat, "", "", "", "", ""],
            ["Testwerk", "Gurt Alpha", "DB-0010", "2019", "", "", "", "", ""]])))
        z = nach_nr(b)
        assert codes(z[2]) == ["datum"] and z[3]["status"] == "uebernehmbar"
        assert codes(z[11]) == ["datum"] and "MM/JJJJ" in z[11]["gruende"][0]["text"]
        assert z[4]["status"] == "uebernehmbar" and any("Kaufdatum „Frühjahr 2018“" in h for h in z[4]["hinweise"])
        assert z[5]["status"] == "uebernehmbar" and z[5]["faellig"] == T.isoformat()
        assert codes(z[6]) == ["intervall"] and codes(z[7]) == ["datum"]
        assert z[8]["status"] == "uebernehmbar" and any("nur bei textilem Material" in h for h in z[8]["hinweise"])
        assert z[9]["status"] == "uebernehmbar" and any("Zukunft" in h for h in z[9]["hinweise"])
        assert z[10]["status"] == "uebernehmbar" and not any("Herstelldatum" in h for h in z[10]["hinweise"])
        r = ok(uebernehmen(t.w, t.lager, csv_bytes([kopf, ["Testwerk", "HMS Beta", "DB-0002", "Baujahr 2019", "Frühjahr 2018",
                                                           "", "Frühjahr 23", "", ""]])), 201)
        iid = r["zeilen"][0]["artikel"]["id"]
        a = t.w.get(f"/api/material/artikel/{iid}").json["artikel"]
        assert a["notes"] == ("Herstelldatum laut Liste: Baujahr 2019\nKaufdatum laut Liste: Frühjahr 2018\n"
                              "Zuletzt geprüft laut Liste: Frühjahr 23"), a["notes"]
        assert a["manufactured_on"] is None and a["purchased_on"] is None and a["next_due_on"] == T.isoformat()
    print("Daten und Blockieren bestanden.")


# --- 8. Fälligkeit --------------------------------------------------------------------------------------------------

def test_faelligkeit():
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        T = _heute()
        grenze = plus(T, 12)
        kopf = ["Hersteller", "Modell", "Seriennummer", "Herstelldatum", "Zuletzt geprüft", "Nächste Prüfung"]
        monat = plus(T, 5)
        zeilen_ = [kopf,
                   ["Testwerk", "HMS Beta", "FA-0001", "", "", tag(plus(T, 3))],                       # 2: N
                   ["Testwerk", "HMS Beta", "FA-0002", "", "", tag(plus(T, 14))],                      # 3: N > T + iv
                   ["Testwerk", "HMS Beta", "FA-0003", "", "", tag(T - timedelta(days=10))],           # 4: N vorbei
                   ["Testwerk", "HMS Beta", "FA-0004", "", tag(plus(T, -3)), ""],                      # 5: nur L
                   ["Testwerk", "HMS Beta", "FA-0005", "", tag(plus(T, -14)), ""],                     # 6: L + iv vorbei
                   ["Testwerk", "HMS Beta", "FA-0006", "", tag(T + timedelta(days=5)), ""],            # 7: L Zukunft
                   ["Testwerk", "HMS Beta", "FA-0007", "05/2020", "01.01.2019", ""],                  # 8: L vor Herst.
                   ["Testwerk", "HMS Beta", "FA-0008", "", tag(plus(T, -9)), tag(plus(T, 6))],         # 9: beide
                   ["Testwerk", "HMS Beta", "FA-0009", "", "", ""],                                    # 10: keins
                   ["Testwerk", "HMS Beta", "FA-0010", "", "", f"{monat.month:02d}/{monat.year}"],     # 11: Monat
                   ["Testwerk", "Gurt Alpha", "FA-0011", f"01/{T.year - 11}", "", ""],                 # 12: Lebensende
                   ["Testwerk", "HMS Beta", "FA-0012", "", "", tag(T + timedelta(days=200))]]          # 13: N in 200 T.
        b = ok(probelauf(t.w, t.lager, csv_bytes(zeilen_)))
        z = nach_nr(b)
        assert all(x["status"] == "uebernehmbar" for x in b["zeilen"]), [(x["nr"], x["gruende"]) for x in b["zeilen"]]
        assert z[2]["faellig"] == plus(T, 3).isoformat()
        assert z[3]["faellig"] == grenze.isoformat() and any("mehr als ein Intervall" in h for h in z[3]["hinweise"])
        assert z[4]["faellig"] == (T - timedelta(days=10)).isoformat()
        assert z[5]["faellig"] == plus(plus(T, -3), 12).isoformat()
        assert z[6]["faellig"] == plus(plus(T, -14), 12).isoformat()
        assert z[7]["faellig"] == T.isoformat() and any("Zukunft" in h for h in z[7]["hinweise"])
        assert z[8]["faellig"] == T.isoformat() and any("vor dem Herstelldatum" in h for h in z[8]["hinweise"])
        assert z[9]["faellig"] == plus(T, 3).isoformat() and any("vorgezogen" in h for h in z[9]["hinweise"])
        assert z[10]["faellig"] == T.isoformat()
        assert z[11]["faellig"] == date(monat.year, monat.month, 1).isoformat()
        assert any("Monatsangabe" in h for h in z[11]["hinweise"])
        assert any("Lebensdauer erreicht" in h for h in z[12]["hinweise"])
        assert b["faelligkeit"] == {"laut_liste_ueberfaellig": 2, "ohne_pruefdatum": 4, "lebensdauer_erreicht": 1}, \
            b["faelligkeit"]
        r = ok(uebernehmen(t.w, t.lager, csv_bytes(zeilen_)), 201)
        z = nach_nr(r)
        ids = {n: z[n]["artikel"]["id"] for n in z}
        assert not zeilen(t.app, "SELECT 1 FROM mat_item WHERE last_inspected_on IS NOT NULL")
        a = t.w.get(f"/api/material/artikel/{ids[10]}").json["artikel"]
        assert a["status"] == "neu" and a["pruefarten"] == {"arten": ["eingang"], "vorschlag": "eingang"}, a["status"]
        assert t.w.get(f"/api/material/artikel/{ids[12]}").json["artikel"]["status"] == "lebensdauer"
        # Zusammenspiel mit den Erinnerungen (L2): fällig heute und überfällig – „neu“; in 200 Tagen nicht.
        with t.app.app_context():
            from app.material.services import erinnerung
            eintraege, _ = erinnerung.artikel_kandidaten(T)
        kand = {e.kennung: e for e in eintraege if e.art == "artikel"}
        assert ids[10] in kand and ids[4] in kand and kand[ids[10]].daten["neu"] and kand[ids[4]].daten["neu"]
        assert ids[13] not in kand
    print("Fälligkeit bestanden.")


# --- 9. Produktabgleich ---------------------------------------------------------------------------------------------

def test_produktabgleich():
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        kar = t.ids["karabiner"]
        vv = produkt(t.mv, kar, "Kopfwerk", "Vertex Vent", "VV-1")
        vv_hv = produkt(t.mv, kar, "Kopfwerk", "Vertex Vent | Hi-Viz", "VV-2")
        sq = produkt(t.mv, kar, "Baumwerk", "Sequoia | SRT", "SQ-1")
        demo = produkt(t.mv, kar, "Demowerk", "Demo Gerät", "DE-1")
        sql(t.app, "UPDATE mat_product SET demo = 1 WHERE id = ?", (demo,))
        kopf = ["Hersteller", "Modell", "Artikelnummer", "Seriennummer"]
        b = ok(probelauf(t.w, t.lager, csv_bytes([kopf,
            ["testwerk", "hms-beta", "HB 7", "PA-0001"],            # 2: genau (Schreibweise egal)
            ["Test-Werk", "HMS B", "HB-7", "PA-0002"],              # 3: Artikelnummer, Hersteller anders geschrieben
            ["Fremdwerk", "HMS Beta", "HB-7", "PA-0003"],           # 4: Artikelnummer bei anderem Hersteller → Vorschlag
            ["Testwerk", "HMS Beta", "", "PA-0004"],                # 5: Modellname, genau ein Produkt
            ["Kopfwerk", "Vertex Vent", "", "PA-0005"],             # 6: zwei Varianten → mehrdeutig
            ["Baumwerk", "Sequoia", "", "PA-0006"],                 # 7: nur „Sequoia | SRT“ → Modell
            ["Testwerk", "HMS Beta", "HB-99", "PA-0007"],           # 8: andere Artikelnummer → Vorschlag
            ["Testwerk", "Gurt Alph", "", "PA-0008"],               # 9: unscharf → Vorschlag
            ["Nirgendwerk", "Phantom", "", "PA-0009"],              # 10: unbekannt
            ["", "HMS Beta", "", "PA-0010"],                        # 11: Hersteller fehlt
            ["Demowerk", "Demo Gerät", "DE-1", "PA-0011"]])))       # 12: DEMO-Produkt im echten Ordner
        z = nach_nr(b)
        for n, pid in ((2, t.ids["hms"]), (3, t.ids["hms"]), (5, t.ids["hms"]), (7, sq)):
            assert z[n]["status"] == "uebernehmbar" and z[n]["produkt"]["id"] == pid, (n, z[n])
        for n in (4, 8, 9):
            assert codes(z[n]) == ["produkt_vorschlag"], (n, z[n])
        assert codes(z[6]) == ["produkt_mehrdeutig"] and codes(z[10]) == ["produkt_unbekannt"]
        assert codes(z[11]) == ["produkt_fehlt"] and codes(z[12]) == ["produkt_unbekannt"], (z[11], z[12])
        offen = {p["angabe"]: p for p in b["produkte_offen"]}
        assert offen["kopfwerk|vertexvent|"]["art"] == "mehrdeutig" and offen["nirgendwerk|phantom|"]["art"] == "unbekannt"
        # Vorschläge der Nacharbeit: Varianten beide, nie ein DEMO-Produkt im echten Ordner.
        ok(uebernehmen(t.w, t.lager, csv_bytes([kopf, ["Kopfwerk", "Vertex Vent", "", "PA-0005"],
                                                ["Demowerk", "Demo Geraet", "", "PA-0012"]])), 201)
        n = ok(t.w.get(f"/api/material/ordner/{t.lager}/nacharbeit"))
        karten = {g_["angabe"]: g_ for g_ in n["gruppen"]}
        assert {p["id"] for p in karten["kopfwerk|vertexvent|"]["vorschlaege"]} == {vv, vv_hv}
        assert all(p["id"] != demo for g_ in n["gruppen"] for p in g_["vorschlaege"])
        # Im DEMO-Ordner passt das DEMO-Produkt.
        gdemo, _ = gliederung(t.adm, "DEMO Platzhalter")
        sql(t.app, "UPDATE mat_folder SET demo = 1 WHERE id = ?", (gdemo,))
        b = ok(probelauf(t.w, gdemo, csv_bytes([kopf, ["Demowerk", "Demo Gerät", "DE-1", "PA-0013"]])))
        assert b["zeilen"][0]["status"] == "uebernehmbar" and b["zeilen"][0]["produkt"]["id"] == demo
    print("Produktabgleich bestanden.")


# --- 10. Probelauf schreibt nichts ------------------------------------------------------------------------------------

def test_probelauf_schreibt_nichts():
    with ohne_post(), tempfile.TemporaryDirectory() as tmp:
        t = zwei_gliederungen(tmp)
        for s in ("FREMD-0002", "FREMD-0003"):
            neuer_artikel(t.wart_b, t.ids["hms"], t.lager_b, serial_no=s)
        neuer_artikel(t.wart_a, t.ids["hms"], t.lager_a, serial_no="EIGEN-0001")
        daten = csv_bytes([KOPF + ["Anzahl"], ["Testwerk", "HMS Beta", "NEU-0001", ""],
                           ["Testwerk", "HMS Beta", "EIGEN-0001", ""], ["Testwerk", "HMS Beta", "TEST-B-0001", ""],
                           ["Testwerk", "HMS Beta", "FREMD-0002", ""], ["Testwerk", "HMS Beta", "FREMD-0003", ""],
                           ["Nirgendwerk", "Phantom", "NEU-0002", ""], ["Testwerk", "HMS Beta", "NEU-0003", "4"],
                           ["Testwerk", "HMS Beta", "", ""]])
        vorher = stand_aller_tabellen(t.app)
        b = ok(probelauf(t.wart_a, t.lager_a, daten, "Geheimer Dateiname.csv"))
        assert {x["status"] for x in b["zeilen"]} == {"uebernehmbar", "im_bestand", "nacharbeit"}
        nachher = stand_aller_tabellen(t.app)
        assert nachher.pop("mat_global_search") == vorher.pop("mat_global_search") + 1
        assert nachher == vorher
        p = zeilen(t.app, "SELECT * FROM mat_global_search ORDER BY id DESC LIMIT 1")[0]
        fremde = {r["id"] for r in zeilen(t.app, "SELECT id FROM mat_item WHERE serial_no IN ('TEST-B-0001', "
                                                 "'FREMD-0002', 'FREMD-0003')")}
        assert p["art"] == "dublette" and "Geheimer" not in p["anfrage"] and set(json.loads(p["item_ids"])) == fremde
        assert p["anfrage"] == "Bestandsimport: 3 Seriennummern" and p["treffer"] == 3
    print("Probelauf schreibt nichts bestanden.")


# --- 11. Übernehmen -------------------------------------------------------------------------------------------------

def test_uebernehmen():
    from app.material.services import Abgelehnt, bestand
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        fach = ordner_neu(t.w, t.lager, "Fach 1")
        kopf = ["Hersteller", "Modell", "Seriennummer", "Inventarnummer", "Ordner", "Anzahl"]
        daten = csv_bytes([kopf,
                           ["Testwerk", "HMS Beta", "UE-0001", "OG-17", "", ""],
                           ["Testwerk", "HMS Beta", "UE-0002", "", "Fach 1", ""],
                           ["Nirgendwerk", "Phantom", "UE-0003", "", "", ""],
                           ["Testwerk", "HMS Beta", "", "OG-18", "", ""],
                           ["Testwerk", "HMS Beta", "UE-0004", "", "", "5"],
                           ["Testwerk", "HMS Beta", "UE-0005", "", "Garage", ""],
                           ["Testwerk", "HMS Beta", "UE-0006", "", "", "1"]])
        zaehler = zeilen(t.app, "SELECT wert FROM mat_meta WHERE schluessel = 'inventar_zaehler'")[0]["wert"]
        # Eine Transaktion: scheitert das Schreiben beim zweiten Artikel, bleibt nichts.
        alt, aufrufe = bestand.artikel_schreiben, []

        def zweiter_scheitert(plan, zusatz=None):
            aufrufe.append(1)
            if len(aufrufe) == 2:
                raise Abgelehnt("Testfehler beim Schreiben", 409)
            return alt(plan, zusatz)
        bestand.artikel_schreiben = zweiter_scheitert
        try:
            r = uebernehmen(t.w, t.lager, daten)
        finally:
            bestand.artikel_schreiben = alt
        assert r.status_code == 409 and "Testfehler" in r.json["error"], r.get_json()
        assert not zeilen(t.app, "SELECT 1 FROM mat_item") and not zeilen(t.app, "SELECT 1 FROM mat_import")
        assert zeilen(t.app, "SELECT wert FROM mat_meta WHERE schluessel = 'inventar_zaehler'")[0]["wert"] == zaehler
        r = ok(uebernehmen(t.w, t.lager, daten, "Bestand Musterstadt.csv"), 201)
        z = nach_nr(r)
        assert r["import"] == {"id": r["import"]["id"], "angelegt": 4, "im_bestand": 0, "nacharbeit": 3, "ersetzt": 0}
        assert [z[n]["status"] for n in range(2, 9)] == ["angelegt", "angelegt", "nacharbeit", "angelegt", "nacharbeit",
                                                         "nacharbeit", "angelegt"]
        nummern = [z[n]["artikel"]["inventory_no"] for n in (2, 3, 5, 8)]
        assert nummern == ["M-000001", "M-000002", "M-000003", "M-000004"], nummern
        assert z[2]["bisherige_nummer"] == "OG-17" and z[5]["seriennummer"] == "STR-000003"
        assert r["ordner_unbekannt"] == [{"wert": "Garage", "zeilen": 1}]
        items = {i["id"]: i for i in zeilen(t.app, "SELECT * FROM mat_item")}
        assert items[z[3]["artikel"]["id"]]["folder_id"] == fach and items[z[2]["artikel"]["id"]]["folder_id"] == t.lager
        assert items[z[2]["artikel"]["id"]]["notes"] == "Bisherige Inventarnummer: OG-17"
        imp = zeilen(t.app, "SELECT * FROM mat_import")[0]
        assert (imp["file_name"], imp["file_kind"], imp["rows_total"], imp["rows_created"], imp["rows_rework"]) == (
            "Bestand Musterstadt.csv", "csv", 7, 4, 3)
        assert re.fullmatch(r"[0-9a-f]{64}", imp["file_sha256"]) and imp["created_by_name"] == "wart"
        j = zeilen(t.app, "SELECT * FROM mat_audit WHERE entity = 'item' AND action = 'angelegt' AND entity_id = ?",
                   (z[2]["artikel"]["id"],))[0]
        nachher = json.loads(j["after_json"])
        assert j["user_name"] == "wart" and nachher["bestandsimport"] == imp["id"] and nachher["zeile"] == 2
        assert nachher["liste_inventarnummer"] == "OG-17" and "datei" not in nachher
        assert "Musterstadt.csv" not in json.dumps(zeilen(t.app, "SELECT * FROM mat_audit WHERE entity = 'item'"))
        j = zeilen(t.app, "SELECT * FROM mat_audit WHERE entity = 'folder' AND action = 'bestandsimport'")[0]
        assert json.loads(j["after_json"]) == {"bestandsimport": imp["id"], "datei": "Bestand Musterstadt.csv",
                                               "zeilen": 7, "artikel_angelegt": 4, "schon_im_bestand": 0,
                                               "nacharbeit": 3, "ersetzt": 0}
        offen = offene_zeilen(t.app, t.lager)
        assert [(o["line_no"], o["reason_code"]) for o in offen] == [(4, "produkt_unbekannt"), (6, "sammelzeile"),
                                                                     (7, "ordner")]
        assert json.loads(offen[2]["values_json"]) == {"hersteller": "Testwerk", "modell": "HMS Beta",
                                                       "seriennummer": "UE-0005", "ordner": "Garage"}
        assert "Garage" in offen[2]["reason"] and offen[0]["offer_key"] == "nirgendwerk|phantom|"
    print("Übernehmen bestanden.")


# --- 12. Wiederholung -----------------------------------------------------------------------------------------------

def test_wiederholung():
    from app.material.services import bestand, bestandsimport
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        daten = csv_bytes([KOPF, ["Testwerk", "HMS Beta", "WH-0001"], ["Testwerk", "HMS Beta", "WH-0002"],
                           ["Testwerk", "HMS Beta", ""]])
        ok(uebernehmen(t.w, t.lager, daten), 201)
        for c in (t.w, t.mv):
            r = uebernehmen(c, t.lager, daten)
            assert r.status_code == 409 and "schon in diesen Ordner übernommen" in r.json["error"], r.get_json()
        assert len(zeilen(t.app, "SELECT 1 FROM mat_item")) == 3
        alt = (datetime.now(timezone.utc) - timedelta(minutes=11)).replace(microsecond=0).isoformat()
        sql(t.app, "UPDATE mat_import SET created_at = ?", (alt,))
        b = ok(probelauf(t.w, t.lager, daten))
        assert b["schon_importiert"][0]["angelegt"] == 3 and "noch einmal angelegt" in b["warnungen"][0]
        assert [x["status"] for x in b["zeilen"]] == ["im_bestand", "im_bestand", "uebernehmbar"]
        r = ok(uebernehmen(t.w, t.lager, daten), 201)
        assert r["import"]["angelegt"] == 1 and r["import"]["im_bestand"] == 2       # nur die Zeile ohne Seriennummer
        # Zwei Übernahmen verschiedener Dateien mit derselben Nummer: unter der Sperre wird noch einmal geprüft – ein
        # Artikel, der zwischen Prüfung und Schreiben entsteht, wird „schon im Bestand“.
        alt_baum = bestandsimport._baum_unveraendert

        def dazwischen(fid, ordner_ids):
            bestand.artikel_anlegen({"product_id": t.ids["hms"], "folder_id": t.lager, "serial_no": "WH-0009"})
            return alt_baum(fid, ordner_ids)
        bestandsimport._baum_unveraendert = dazwischen
        try:
            r = ok(uebernehmen(t.w, t.lager, csv_bytes([KOPF, ["Testwerk", "HMS Beta", "WH-0009"]])), 201)
        finally:
            bestandsimport._baum_unveraendert = alt_baum
        assert r["zeilen"][0]["status"] == "im_bestand" and r["import"]["angelegt"] == 0
        assert len(zeilen(t.app, "SELECT 1 FROM mat_item WHERE serial_no = 'WH-0009'")) == 1
    print("Wiederholung bestanden.")


# --- 13. Nacharbeits-CSV im Rundweg ---------------------------------------------------------------------------------

def _csv_lesen(daten):
    return list(csv.reader(io.StringIO(daten.decode("utf-8-sig")), delimiter=";"))


def test_nacharbeit_csv_rundweg():
    from app.material.services import Abgelehnt, bestand
    with ohne_post(), tempfile.TemporaryDirectory() as tmp:
        t = zwei_gliederungen(tmp)
        app, w = t.app, t.wart_a
        kopf = ["Hersteller", "Modell", "Seriennummer", "Herstelldatum", "Prüfintervall", "Anzahl", "Bemerkung"]
        ok(uebernehmen(w, t.lager_a, csv_bytes([kopf,
            ["Testwerk", "Gurt Alpha", "RW-0001", "", "", "", "=SUMME(A1)"],      # textil ohne Herstelldatum
            ["Testwerk", "Gurt Alpha", "", "", "", "", ""],                       # dasselbe ohne Seriennummer
            ["Testwerk", "HMS Beta", "RW-0003", "", "jährlich?", "", ""],         # Intervall unlesbar
            ["Nirgendwerk", "Phantom", "RW-0004", "", "", "3", ""],               # Produkt und Sammelzeile
            ["Testwerk", "HMS Beta", "RW-0005", "", "", "", ""]])), 201)          # übernommen
        ok(uebernehmen(t.wart_b, t.lager_b, csv_bytes([KOPF, ["Nirgendwerk", "Phantom", "RB-0001"]])), 201)
        b_zeile = offene_zeilen(app, t.lager_b)[0]
        offen = offene_zeilen(app, t.lager_a)
        assert len(offen) == 4
        ok(w.post(f"/api/material/ordner/{t.lager_a}/nacharbeit/produkt", json={"zeilen": [offen[3]["id"]],
                                                                                  "product_id": t.ids["hms"]}))
        r = w.get(f"/material/ordner/{t.lager_a}/nacharbeit.csv")
        assert r.status_code == 200 and r.data.startswith(b"\xef\xbb\xbf") and "no-store" in r.headers["Cache-Control"]
        tabelle_ = _csv_lesen(r.data)
        assert tabelle_[0] == ["Nacharbeit-Nr.", "Grund", "Hersteller", "Modell", "Artikelnummer", "Seriennummer",
                               "Bisherige Inventarnummer", "Ordner", "Herstelldatum", "Kaufdatum", "Erstgebrauch",
                               "Zuletzt geprüft", "Nächste Prüfung", "Prüfintervall (Monate)", "Lebensende", "Anzahl",
                               "Bemerkung", "Zeile"], tabelle_[0]
        assert [int(z[0]) for z in tabelle_[1:]] == [o["id"] for o in offen] and tabelle_[1][16] == "'=SUMME(A1)"
        # Korrigiert: Herstelldatum ergänzt, Intervall richtig; dazu eine Nacharbeit-Nr. aus B. Die Zeile mit
        # zugeordnetem Produkt bleibt eine Sammelzeile (Anzahl 3).
        korr = [tabelle_[0]]
        for z in tabelle_[1:]:
            z = list(z)
            if z[3] == "Gurt Alpha":
                z[8] = "03/2021"
            if z[5] == "RW-0003":
                z[13] = "12"
            korr.append(z)
        korr.append([str(b_zeile["id"]), "", "Testwerk", "HMS Beta", "", "RW-0009"] + [""] * 12)
        korr_daten = csv_bytes(korr)
        b = ok(probelauf(w, t.lager_a, korr_daten, "korrigiert.csv"))
        assert b["nacharbeit_csv"] and not b["fehlt"] and any("gehören nicht zur Nacharbeit" in x for x in b["warnungen"])
        # Scheitert das Schreiben, bleiben auch die alten Zeilen offen.
        alt = bestand.artikel_schreiben
        bestand.artikel_schreiben = lambda *a, **k: (_ for _ in ()).throw(Abgelehnt("Testfehler", 409))
        try:
            assert uebernehmen(w, t.lager_a, korr_daten, "korrigiert.csv").status_code == 409
        finally:
            bestand.artikel_schreiben = alt
        assert len(offene_zeilen(app, t.lager_a)) == 4
        r = ok(uebernehmen(w, t.lager_a, korr_daten, "korrigiert.csv"), 201)
        assert r["import"]["angelegt"] == 4 and r["import"]["ersetzt"] == 4 and r["import"]["nacharbeit"] == 1, r["import"]
        alte = zeilen(app, f"SELECT * FROM mat_import_row WHERE id IN ({','.join(str(o['id']) for o in offen)})")
        assert {a["status"] for a in alte} == {"ersetzt"} and {a["values_json"] for a in alte} == {"{}"}
        assert all(a["reason"].startswith(f"neu eingelesen mit Bestandsimport Nr. {r['import']['id']}") for a in alte)
        neu_offen = offene_zeilen(app, t.lager_a)
        assert len(neu_offen) == 1 and neu_offen[0]["product_id"] == t.ids["hms"], neu_offen
        assert zeilen(app, "SELECT status, product_id FROM mat_import_row WHERE id = ?", (b_zeile["id"],))[0] == {
            "status": "offen", "product_id": None}
        anzahl = len(zeilen(app, "SELECT 1 FROM mat_item"))
        # Dieselbe Datei neu gespeichert (andere Prüfsumme) und nach 11 Minuten noch einmal: alles erledigt.
        sql(app, "UPDATE mat_import SET created_at = ?",
            ((datetime.now(timezone.utc) - timedelta(minutes=11)).replace(microsecond=0).isoformat(),))
        b = ok(probelauf(w, t.lager_a, korr_daten + b"\r\n", "korrigiert2.csv"))
        assert [z["status"] for z in b["zeilen"]] == ["erledigt"] * 4 + ["im_bestand"], b["zeilen"]
        assert all(z["hinweise"][0].startswith("schon erledigt (neu eingelesen am") for z in b["zeilen"][:4])
        r = uebernehmen(w, t.lager_a, korr_daten + b"\r\n", "korrigiert2.csv")
        assert r.status_code == 409 and "Nichts zu übernehmen" in r.json["error"]
        assert len(zeilen(app, "SELECT 1 FROM mat_item")) == anzahl            # kein Artikel doppelt, auch kein STR
    print("Nacharbeits-CSV im Rundweg bestanden.")


# --- 14. Varianten zeilenweise --------------------------------------------------------------------------------------

def test_varianten_zeilenweise():
    with ohne_post(), tempfile.TemporaryDirectory() as tmp:
        t = zwei_gliederungen(tmp)
        app, w, kar = t.app, t.wart_a, t.ids["karabiner"]
        mv = t.k["katalog"]
        var = {g_: produkt(mv, kar, "Baumwerk", f"Sequoia | {g_}", f"SQ-{g_}") for g_ in ("S", "M", "L")}
        kopf = ["Hersteller", "Modell", "Größe", "Seriennummer"]
        ok(uebernehmen(w, t.lager_a, csv_bytes([kopf] + [["Baumwerk", "Sequoia", g_, f"SV-000{n}"]
                                                         for n, g_ in enumerate(("S", "M", "L"), 1)])), 201)
        offen = offene_zeilen(app, t.lager_a)
        assert [o["reason_code"] for o in offen] == ["produkt_mehrdeutig"] * 3
        n = ok(w.get(f"/api/material/ordner/{t.lager_a}/nacharbeit"))
        assert n["gruppen"][0]["art"] == "mehrdeutig" and {p["id"] for p in n["gruppen"][0]["vorschlaege"]} == set(
            var.values())
        for o, g_ in zip(offen, ("S", "M", "L")):
            assert ok(w.post(f"/api/material/ordner/{t.lager_a}/nacharbeit/produkt",
                             json={"zeilen": [o["id"]], "product_id": var[g_]}))["gesetzt"] == 1
        r = ok(w.post(f"/api/material/ordner/{t.lager_a}/nacharbeit/pruefen", json={"zeilen": "alle",
                                                                                     "uebernehmen": True}))
        assert r["zahlen"]["angelegt"] == 3, r["zahlen"]
        assert {i["product_id"] for i in zeilen(app, "SELECT product_id FROM mat_item WHERE serial_no LIKE 'SV-%'")} \
            == set(var.values())
        # Gleiche Seriennummer bei einer anderen Variante desselben Herstellers – sichtbar: Grund mit Inventarnummer;
        # fremd: kein Grund (kein zweiter Weg zum Abzählen).
        eigen = neuer_artikel(w, var["M"], t.lager_a, serial_no="SQ-0001")
        neuer_artikel(t.wart_b, var["M"], t.lager_b, serial_no="SQ-0002")
        b = ok(probelauf(w, t.lager_a, csv_bytes([["Hersteller", "Modell", "Artikelnummer", "Seriennummer"],
                                                  ["Baumwerk", "Sequoia | S", "SQ-S", "SQ-0001"],
                                                  ["Baumwerk", "Sequoia | S", "SQ-S", "SQ-0002"]])))
        z = nach_nr(b)
        assert codes(z[2]) == ["serie_variante"] and eigen["inventory_no"] in z[2]["gruende"][0]["text"]
        assert z[3]["status"] == "uebernehmbar" and "Nachbarstadt" not in json.dumps(b)
    print("Varianten zeilenweise bestanden.")


# --- 15. Sammelzeilen -----------------------------------------------------------------------------------------------

def test_sammelzeilen():
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        r = ok(uebernehmen(t.w, t.lager, csv_bytes([KOPF + ["Anzahl"], ["Testwerk", "HMS Beta", "SZ-0001", "10"],
                                                    ["Testwerk", "HMS Beta", "SZ-0002", "1"],
                                                    ["Testwerk", "HMS Beta", "SZ-0003", ""],
                                                    ["Testwerk", "HMS Beta", "SZ-0004", "ca. 5"],
                                                    ["Testwerk", "HMS Beta", "SZ-0005", "0"]])), 201)
        z = nach_nr(r)
        assert [codes(z[n]) for n in (2, 5, 6)] == [["sammelzeile"]] * 3 and "Anzahl 10" in z[2]["gruende"][0]["text"]
        assert z[3]["status"] == z[4]["status"] == "angelegt"
        assert [i["serial_no"] for i in zeilen(t.app, "SELECT serial_no FROM mat_item ORDER BY id")] == ["SZ-0002",
                                                                                                         "SZ-0003"]
    print("Sammelzeilen bestanden.")


# --- 16. Ordnerspalte -----------------------------------------------------------------------------------------------

def test_ordnerspalte():
    from app.material.services import bestandsimport
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        w, gid = t.w, t.gid
        regal = ordner_neu(w, t.lager, "Regal 1")
        fach = ordner_neu(w, regal, "Fach 2")
        a, b_ = ordner_neu(w, gid, "Halle A"), ordner_neu(w, gid, "Halle B")
        fa, fb = ordner_neu(w, a, "Fach"), ordner_neu(w, b_, "Fach")
        ordner_neu(w, gid, "Bühne/Seite")
        buehne = ordner_neu(w, gid, "Bühne")
        ordner_neu(w, buehne, "Seite")
        werte = ["Lager Musterstadt / Regal 1 / Fach 2", "lager musterstadt/REGAL 1/fach 2", "Halle A/Fach", "Halle B / Fach",
                 "Garage", "Garage", "Gesperrt", "Bühne / Seite", ""]
        bericht = ok(probelauf(w, gid, csv_bytes([KOPF + ["Ordner"]] + [["Testwerk", "HMS Beta", f"OS-000{n}", v]
                                                                        for n, v in enumerate(werte)])))
        z = nach_nr(bericht)
        ordner_ids = {r["id"]: r for r in zeilen(t.app, "SELECT id, name FROM mat_folder")}
        assert z[2]["status"] == z[3]["status"] == "uebernehmbar" and z[2]["ordner"].endswith("Regal 1 / Fach 2")
        assert z[4]["ordner"].endswith("Halle A / Fach") and z[5]["ordner"].endswith("Halle B / Fach")
        assert codes(z[6]) == codes(z[7]) == codes(z[8]) == codes(z[9]) == ["ordner"]
        assert "mehrdeutig" in z[9]["gruende"][0]["text"] and "Gesperrt" in z[8]["gruende"][0]["text"]
        assert bericht["ordner_unbekannt"] == [{"wert": "Garage", "zeilen": 2}]
        assert z[10]["ordner"] == "OG Musterstadt" and fach in ordner_ids and fa != fb
        r = probelauf(w, t.system["gesperrt"], csv_bytes([KOPF, ["Testwerk", "HMS Beta", "OS-0100"]]))
        assert r.status_code == 409 and "wird nichts importiert" in r.json["error"]
        # Zwischen Prüfung und Schreibsperre wandert ein Unterordner in eine andere Gliederung: 409, nichts geschrieben.
        andere, _ = gliederung(t.adm, "OG Anderswo")
        alt = bestandsimport._pruefen

        def verschieben_danach(*a, **k):
            b = alt(*a, **k)
            sql(t.app, "UPDATE mat_folder SET parent_id = ? WHERE id = ?", (andere, a_halle))
            return b
        a_halle = a
        bestandsimport._pruefen = verschieben_danach
        try:
            r = uebernehmen(w, gid, csv_bytes([KOPF + ["Ordner"], ["Testwerk", "HMS Beta", "OS-0200", "Halle A / Fach"]]))
        finally:
            bestandsimport._pruefen = alt
        assert r.status_code == 409 and "Ordnerbaum hat sich geändert" in r.json["error"], r.get_json()
        assert not zeilen(t.app, "SELECT 1 FROM mat_item") and not zeilen(t.app, "SELECT 1 FROM mat_import")
    print("Ordnerspalte bestanden.")


# --- 17. Pflichtangaben und Regeln ----------------------------------------------------------------------------------

def test_pflichtangaben_und_regeln():
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        kar = t.ids["karabiner"]
        ohne_iv = produkt(t.mv, kar, "Leerwerk", "Ohne Intervall", "LW-1")
        sql(t.app, "UPDATE mat_product SET interval_months = NULL WHERE id = ?", (ohne_iv,))
        produkt(t.mv, kar, "Langwerk", "Zwei Jahre", "LG-1", interval_months=24)
        T = _heute()
        kopf = ["Hersteller", "Modell", "Seriennummer", "Herstelldatum", "Kaufdatum", "Prüfintervall", "Lebensende",
                "Bemerkung"]
        herg = f"01/{T.year - 2}"
        rechnung = date(T.year + 8, 1, 1)
        b = ok(probelauf(t.w, t.lager, csv_bytes([kopf,
            ["Testwerk", "Gurt Alpha", "PR-0001", "", "", "", "", ""],                                  # 2 pflicht
            ["Leerwerk", "Ohne Intervall", "PR-0002", "", "", "", "", ""],                              # 3 intervall
            ["Testwerk", "HMS Beta", "PR-0003", "", "", "36", "", ""],                                  # 4 länger
            ["Langwerk", "Zwei Jahre", "PR-0004", "", "", "1 Jahr", "", ""],                            # 5 kürzer
            ["Testwerk", "Gurt Alpha", "PR-0005", herg, "", "", tag(rechnung - timedelta(days=400)), ""],  # 6 früher
            ["Testwerk", "Gurt Alpha", "PR-0006", herg, "", "", tag(rechnung + timedelta(days=30)), ""],  # 7 später
            ["Testwerk", "HMS Beta", "PR-0007", "", "", "", tag(rechnung), ""],                          # 8 Metall
            ["Testwerk", "HMS Beta", "PR-0008", "05/2020", "01/2019", "", "", ""],                      # 9 Kauf vor Herst.
            ["Testwerk", "HMS Beta", "PR-0009", "", "", "", "", "x" * 4500]])))                         # 10 lang
        z = nach_nr(b)
        assert codes(z[2]) == ["pflicht"] and codes(z[3]) == ["intervall"] and "Kein Prüfintervall" in z[3]["gruende"][0][
            "text"]
        assert z[4]["status"] == "uebernehmbar" and any("nicht kürzer" in h for h in z[4]["hinweise"])
        assert z[5]["status"] == z[6]["status"] == z[7]["status"] == z[8]["status"] == "uebernehmbar"
        assert any("vorgezogen" in h for h in z[6]["hinweise"]) and any("später" in h for h in z[7]["hinweise"])
        assert any("nur bei textilem Material" in h for h in z[8]["hinweise"])
        assert codes(z[9]) == ["regel"] and z[9]["gruende"][0]["text"] == "Kaufdatum liegt vor dem Herstelldatum."
        assert any("4000 Zeichen gekürzt" in h for h in z[10]["hinweise"])
        r = ok(uebernehmen(t.w, t.lager, csv_bytes([kopf,
            ["Testwerk", "HMS Beta", "PR-0003", "", "", "36", "", ""],
            ["Langwerk", "Zwei Jahre", "PR-0004", "", "", "1 Jahr", "", ""],
            ["Testwerk", "Gurt Alpha", "PR-0005", herg, "", "", tag(rechnung - timedelta(days=400)), ""],
            ["Testwerk", "HMS Beta", "PR-0009", "", "", "", "", "x" * 4500]])), 201)
        i = {x["serial_no"]: x for x in zeilen(t.app, "SELECT * FROM mat_item")}
        assert i["PR-0003"]["interval_months"] is None and i["PR-0004"]["interval_months"] == 12
        assert i["PR-0005"]["retire_by_manual"] == (rechnung - timedelta(days=400)).isoformat()
        assert len(i["PR-0009"]["notes"]) == 4000 and r["import"]["angelegt"] == 4
    print("Pflichtangaben und Regeln bestanden.")


# --- 18. Rückruf ----------------------------------------------------------------------------------------------------

def test_rueckruf_im_probelauf():
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        with t.app.app_context():
            from app import db
            for titel, pid, serien, zu in (("Platzhalterrückruf Liste", t.ids["hms"], "RR-0001\nrr-9999", None),
                                           ("Platzhalterrückruf alle", t.ids["gurt"], "", None),
                                           ("Platzhalterrückruf beendet", t.ids["hms"], "", "2026-01-01T00:00:00+00:00")):
                db.execute("INSERT INTO mat_recall (product_id, title, serials, created_at, closed_at) VALUES (?,?,?,?,?)",
                           (pid, titel, serien, db.now(), zu))
        kopf = ["Hersteller", "Modell", "Seriennummer", "Herstelldatum"]
        daten = csv_bytes([kopf, ["Testwerk", "HMS Beta", "rr-0001", ""], ["Testwerk", "HMS Beta", "RR-0002", ""],
                           ["Testwerk", "Gurt Alpha", "RR-0003", "03/2024"]])
        b = ok(probelauf(t.w, t.lager, daten))
        z = nach_nr(b)
        assert z[2]["rueckruf"] == "Platzhalterrückruf Liste" and "Rückruf: Platzhalterrückruf Liste" in z[2]["hinweise"]
        assert z[3]["rueckruf"] is None
        assert z[4]["rueckruf"] == "Platzhalterrückruf alle" and any("alle Artikel" in h for h in z[4]["hinweise"])
        assert {x["status"] for x in b["zeilen"]} == {"uebernehmbar"}
        r = ok(uebernehmen(t.w, t.lager, daten), 201)
        a = t.w.get(f"/api/material/artikel/{nach_nr(r)[2]['artikel']['id']}").json["artikel"]
        assert [x["title"] for x in a["rueckrufe"]] == ["Platzhalterrückruf Liste"]
    print("Rückruf im Probelauf bestanden.")


# --- 19. Fremde Dublette und Abzählen -------------------------------------------------------------------------------

def test_fremde_dublette_und_abzaehlen():
    with ohne_post(), tempfile.TemporaryDirectory() as tmp:
        t = zwei_gliederungen(tmp)
        app, w = t.app, t.wart_a
        inv_b = zeilen(app, "SELECT inventory_no FROM mat_item WHERE id = ?", (t.art_b["id"],))[0]["inventory_no"]
        neuer_artikel(t.wart_b, t.ids["hms"], t.lager_b, serial_no="17")
        protokoll = len(zeilen(app, "SELECT 1 FROM mat_global_search"))
        daten = csv_bytes([KOPF, ["Testwerk", "HMS Beta", "TEST-B-0001"]])
        for lauf in (probelauf, uebernehmen):
            r = lauf(w, t.lager_a, daten)
            text = r.get_data(as_text=True)
            b = r.get_json()
            assert codes(b["zeilen"][0]) == ["fremde_dublette"] and "Nachbarstadt" not in text and inv_b not in text
            protokoll += 1
            assert len(zeilen(app, "SELECT 1 FROM mat_global_search")) == protokoll
        assert not zeilen(app, "SELECT 1 FROM mat_item WHERE serial_no = 'TEST-B-0001' AND folder_id = ?", (t.lager_a,))
        # Kurze Nummern: gleich, ob es sie fremd gibt – kein Protokoll, nie angelegt.
        b = ok(probelauf(w, t.lager_a, csv_bytes([KOPF, ["Testwerk", "HMS Beta", "17"], ["Testwerk", "HMS Beta", "18"]])))
        assert [codes(z) for z in b["zeilen"]] == [["serie_von_hand"], ["serie_von_hand"]]
        assert [z["gruende"][0]["text"].replace("18", "17") for z in b["zeilen"]][0] == b["zeilen"][0]["gruende"][0]["text"]
        assert b["zeilen"][1]["gruende"][0]["text"] == b["zeilen"][0]["gruende"][0]["text"].replace("„17“", "„18“")
        assert len(zeilen(app, "SELECT 1 FROM mat_global_search")) == protokoll
        neuer_artikel(w, t.ids["hms"], t.lager_a, serial_no="18")
        b = ok(probelauf(w, t.lager_a, csv_bytes([KOPF, ["Testwerk", "HMS Beta", "18"]])))
        assert b["zeilen"][0]["status"] == "im_bestand"                   # die eigene kurze Nummer erkennt er wieder
    print("Fremde Dublette und Abzählen bestanden.")


# --- 20. Rechte -----------------------------------------------------------------------------------------------------

def test_rechte():
    with ohne_post(), tempfile.TemporaryDirectory() as tmp:
        t = zwei_gliederungen(tmp)
        app, w, k = t.app, t.wart_a, t.k
        daten = csv_bytes([KOPF, ["Nirgendwerk", "Phantom", "RE-0001"], ["Nirgendwerk", "Phantom", "RE-0002"]])
        fach = ordner_neu(w, t.lager_a, "Fach Musterstadt")
        ok(uebernehmen(w, fach, daten), 201)
        ok(uebernehmen(t.wart_b, t.lager_b, daten), 201)
        zeile_fach = offene_zeilen(app, fach)[0]["id"]
        zeile_b = offene_zeilen(app, t.lager_b)[0]["id"]
        mitglied, _ = neues_konto(app, "mitglied@example.org", name="Mitglied Platzhalter")

        def routen(fid):
            basis = f"/api/material/ordner/{fid}"
            return [("POST", f"{basis}/import/probelauf", "datei"), ("POST", f"{basis}/import", "datei"),
                    ("GET", f"{basis}/nacharbeit", None), ("POST", f"{basis}/nacharbeit/produkt", {"zeilen": "alle"}),
                    ("POST", f"{basis}/nacharbeit/pruefen", {"zeilen": "alle"}),
                    ("POST", f"{basis}/nacharbeit/verwerfen", {"zeilen": "alle"}),
                    ("POST", f"{basis}/nacharbeit/melden", {"angabe": "x"}),
                    ("GET", f"/material/ordner/{fid}/import", None), ("GET", f"/material/ordner/{fid}/nacharbeit.csv", None)]

        def anfrage(c, methode, url, inhalt):
            if inhalt == "datei":
                return c.post(url, headers=H, content_type="multipart/form-data",
                              data={"file": (io.BytesIO(daten), "liste.csv")})
            if methode == "GET":
                return c.get(url)
            return c.post(url, json=inhalt)
        vorher = stand_aller_tabellen(app)
        for c, fid, soll in ((k["fuehrung"], t.lager_a, 403), (k["pruefer"], t.lager_a, 403),
                             (mitglied, t.lager_a, 404), (w, t.lager_b, 404), (k["nutzer"], t.lager_a, 403)):
            for methode, url, inhalt in routen(fid):
                r = anfrage(c, methode, url, inhalt)
                assert r.status_code == soll, (methode, url, r.status_code, r.get_json(silent=True))
        assert stand_aller_tabellen(app) == vorher
        # Zeilen aus B im Inhalt einer Anfrage an A und Zeilen eines Imports in einen Unterordner von A: 404.
        for zid in (zeile_b, zeile_fach):
            for teil, inhalt in (("produkt", {"zeilen": [zid], "product_id": t.ids["hms"]}), ("pruefen", {"zeilen": [zid]}),
                                 ("verwerfen", {"zeilen": [zid]})):
                r = w.post(f"/api/material/ordner/{t.lager_a}/nacharbeit/{teil}", json=inhalt)
                assert r.status_code == 404 and r.json["error"] == "Zeile nicht gefunden.", (teil, r.get_json())
        assert stand_aller_tabellen(app) == vorher
        # Der Materialwart darf in A alles, die Material-Verwaltung überall.
        assert w.get(f"/api/material/ordner/{fach}/nacharbeit").status_code == 200
        assert w.get(f"/material/ordner/{fach}/nacharbeit.csv").status_code == 200
        for fid in (fach, t.lager_b):
            assert k["katalog"].get(f"/api/material/ordner/{fid}/nacharbeit").json["zahlen"]["offen"] == 2
            assert k["katalog"].get(f"/material/ordner/{fid}/import").status_code == 200
        ok(probelauf(k["katalog"], t.lager_b, daten))
    print("Rechte bestanden.")


# --- 21. Nacharbeit -------------------------------------------------------------------------------------------------

def test_nacharbeit():
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        app, w = t.app, t.w
        andere = ordner_neu(w, t.gid, "Halle Musterstadt")
        kopf = ["Hersteller", "Modell", "Seriennummer", "Anzahl"]
        ok(uebernehmen(w, t.lager, csv_bytes([kopf, ["Testwerk", "HMS Gamma", "NA-0001", ""],
                                              ["Testwerk", "HMS Gamma", "NA-0002", ""],
                                              ["Testwerk", "HMS Gamma", "NA-0003", ""],
                                              ["Testwerk", "HMS Beta", "NA-0004", "3"]])), 201)
        ok(uebernehmen(w, andere, csv_bytes([kopf, ["Testwerk", "HMS Gamma", "NA-0101", ""]])), 201)
        offen = offene_zeilen(app, t.lager)
        angabe = offen[0]["offer_key"]
        url = f"/api/material/ordner/{t.lager}/nacharbeit"
        n = ok(w.get(url))
        assert n["zahlen"] == {"offen": 4, "produkt": 3, "seriennummer": 0, "daten": 0, "ordner": 0, "sonstige": 1}
        assert n["frist"] and all(z["offen_bis"] for z in n["zeilen"]) and "no-store" in w.get(url).headers["Cache-Control"]
        assert [g_["angabe"] for g_ in n["gruppen"]] == [angabe] and n["gruppen"][0]["offen"] == 3
        assert ok(w.post(f"{url}/produkt", json={"angabe": angabe, "product_id": t.ids["hms"]}))["gesetzt"] == 3
        assert offene_zeilen(app, andere)[0]["product_id"] is None
        assert ok(w.get(url))["gruppen"][0]["produkt"]["id"] == t.ids["hms"]
        assert ok(w.post(f"{url}/produkt", json={"zeilen": [offen[2]["id"]], "product_id": None}))["gesetzt"] == 1
        assert [o["product_id"] for o in offene_zeilen(app, t.lager)] == [t.ids["hms"], t.ids["hms"], None, None]
        # Erneut prüfen ohne Übernehmen schreibt nichts.
        vorher = stand_aller_tabellen(app)
        r = ok(w.post(f"{url}/pruefen", json={"zeilen": "alle"}))
        assert r["zahlen"]["uebernehmbar"] == 2 and r["zahlen"]["nacharbeit"] == 2 and r["zeilen"][0]["id"] == offen[0]["id"]
        assert stand_aller_tabellen(app) == vorher
        # Eine Zeile inzwischen von Hand angelegt → schon im Bestand.
        hand = neuer_artikel(w, t.ids["hms"], t.lager, serial_no="NA-0002")
        r = ok(w.post(f"{url}/pruefen", json={"zeilen": [o["id"] for o in offen[:3]], "uebernehmen": True}))
        assert [z["status"] for z in r["zeilen"]] == ["angelegt", "im_bestand", "nacharbeit"], r["zeilen"]
        assert r["zeilen"][0]["artikel"]["inventory_no"].startswith("M-")
        z = {x["id"]: x for x in zeilen(app, "SELECT * FROM mat_import_row WHERE id IN (?, ?, ?)",
                                         tuple(o["id"] for o in offen[:3]))}
        assert (z[offen[0]["id"]]["status"], z[offen[0]["id"]]["values_json"]) == ("uebernommen", "{}")
        assert (z[offen[1]["id"]]["status"], z[offen[1]["id"]]["item_id"]) == ("im_bestand", hand["id"])
        j = zeilen(app, "SELECT after_json FROM mat_audit WHERE entity = 'item' AND entity_id = ? AND action = 'angelegt'",
                   (r["zeilen"][0]["artikel"]["id"],))[0]
        assert json.loads(j["after_json"])["aus_nacharbeit"] is True
        assert json.loads(zeilen(app, "SELECT after_json FROM mat_audit WHERE action = 'nacharbeit_uebernommen'")[0][
            "after_json"]) == {"artikel_angelegt": 1, "schon_im_bestand": 1}
        r = w.post(f"{url}/pruefen", json={"zeilen": [offen[0]["id"]]})
        assert r.status_code == 409 and "schon erledigt" in r.json["error"]
        assert ok(w.post(f"{url}/verwerfen", json={"zeilen": "alle"}))["verworfen"] == 2
        assert not offene_zeilen(app, t.lager) and w.post(f"{url}/verwerfen", json={"zeilen": "alle"}).status_code == 409
        assert json.loads(zeilen(app, "SELECT after_json FROM mat_audit WHERE action = 'nacharbeit_verworfen'")[0][
            "after_json"]) == {"zeilen": 2}
        # Seiten zu 200, Filter nach Gruppe, die letzten 100 erledigten.
        ok(uebernehmen(w, t.lager, csv_bytes([kopf] + [["Nirgendwerk", "Phantom", f"NB-{n:04d}", ""] for n in range(205)]
                                             + [["Testwerk", "HMS Beta", "NC-0001", "7"]])), 201)
        n = ok(w.get(url))
        assert (n["seite"], n["seiten"], len(n["zeilen"]), n["zahlen"]["offen"]) == (1, 2, 200, 206)
        assert len(ok(w.get(url + "?seite=2"))["zeilen"]) == 6
        assert [z["werte"]["seriennummer"] for z in ok(w.get(url + "?gruppe=sonstige"))["zeilen"]] == ["NC-0001"]
        assert w.get(url + "?gruppe=unsinn").status_code == 400
        ok(w.post(f"{url}/verwerfen", json={"zeilen": "alle"}))
        assert len(ok(w.get(url))["erledigt"]) == 100
    print("Nacharbeit bestanden.")


# --- 22. Produkt melden ---------------------------------------------------------------------------------------------

def test_produkt_melden():
    from app.material.services import meldung
    post = []
    alt = (meldung.HINTERGRUND, meldung.send_mail)
    meldung.HINTERGRUND = lambda f, *a: f(*a)
    meldung.send_mail = lambda an, betreff, inhalt: post.append((an, betreff, inhalt)) or True
    try:
        with tempfile.TemporaryDirectory() as tmp:
            t = aufbau(tmp)
            app, w = t.app, t.w
            kopf = ["Hersteller", "Modell", "Seriennummer"]
            ok(uebernehmen(w, t.lager, csv_bytes([kopf, ["Meldewerk", "Platzhaltergerät", "PM-0001"],
                                                  ["Meldewerk", "Platzhaltergerät", "PM-0002"],
                                                  ["Testwerk", "HMS Beta", "PM-0003", ]]), "Liste Person.csv"), 201)
            offen = offene_zeilen(app, t.lager)
            angabe = offen[0]["offer_key"]
            url = f"/api/material/ordner/{t.lager}/nacharbeit"
            r = w.post(f"{url}/melden", json={"angabe": angabe, "category_id": t.ids["karabiner"]})
            m = ok(r, 201)["meldung"]
            imp = offen[0]["import_id"]
            assert (m["manufacturer"], m["model"]) == ("Meldewerk", "Platzhaltergerät")
            assert m["note"] == f"Aus dem Bestandsimport Nr. {imp}: 2 Artikel warten in der Nacharbeit."
            assert "Person" not in m["note"] and m["ordner"]["id"] == t.lager
            assert {o["request_id"] for o in offene_zeilen(app, t.lager)} == {m["id"]}
            karte = ok(w.get(url))["gruppen"][0]
            assert karte["meldung"]["id"] == m["id"] and karte["meldung"]["status"] == "offen"
            r = w.post(f"{url}/melden", headers=H, content_type="multipart/form-data",
                       data={"angabe": angabe, "manufacturer": "Testwerk", "model": "HMS Beta", "article_no": "HB-7"})
            assert r.status_code == 409, r.get_json()
            assert w.post(f"{url}/melden", json={"angabe": "gibt|es|nicht"}).status_code == 404
            # Die Material-Verwaltung pflegt ein: Die Mail an den Melder führt zur Nacharbeit; erneut prüfen findet
            # das Produkt über die Meldung.
            del post[:]
            r = t.mv.post(f"/api/material/meldungen/{m['id']}/produkt",
                          json={"category_id": t.ids["karabiner"], "manufacturer": "Meldewerk", "model": "Platzhaltergerät",
                                "material_kind": "metall_hardware", "interval_months": 12})
            neu = ok(r, 201)["produkt"]["id"]
            mail = [x for x in post if x[0] == "wart@example.org"]
            assert mail and f"/material/ordner/{t.lager}/import" in mail[0][2] and "Artikel anlegen" not in mail[0][2]
            r = ok(w.post(f"{url}/pruefen", json={"zeilen": "alle", "uebernehmen": True}))
            assert r["zahlen"]["angelegt"] == 2 and {z["produkt"]["id"] for z in r["zeilen"]} == {neu}
            # Abgelehnte Meldung: die Karte zeigt die Antwort.
            ok(uebernehmen(w, t.lager, csv_bytes([kopf, ["Ablehnwerk", "Unsinn", "PM-0010"]])), 201)
            angabe = offene_zeilen(app, t.lager)[0]["offer_key"]
            m2 = ok(w.post(f"{url}/melden", json={"angabe": angabe}), 201)["meldung"]
            ok(t.mv.post(f"/api/material/meldungen/{m2['id']}/ablehnen", json={"antwort": "Kein PSA-Produkt."}))
            karte = ok(w.get(url))["gruppen"][0]
            assert karte["meldung"]["status"] == "abgelehnt" and karte["meldung"]["antwort"] == "Kein PSA-Produkt."
            # Die Material-Verwaltung legt selbst an und ordnet zu.
            pid = produkt(t.mv, t.ids["karabiner"], "Ablehnwerk", "Unsinn", "AW-1")
            assert ok(t.mv.post(f"{url}/produkt", json={"angabe": angabe, "product_id": pid}))["gesetzt"] == 1
            assert offene_zeilen(app, t.lager)[0]["product_id"] == pid
    finally:
        meldung.HINTERGRUND, meldung.send_mail = alt
    print("Produkt melden bestanden.")


# --- 23. Last und Leistung ------------------------------------------------------------------------------------------

def test_last_und_leistung():
    from app import db as db_mod
    from app.material.services import bestandsimport
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        w = t.w
        daten = csv_bytes([KOPF] + [["Testwerk", "HMS Beta", f"LL-{n:05d}"] for n in range(2000)])
        url = f"/api/material/ordner/{t.lager}"
        ok(uebernehmen(w, t.lager, csv_bytes([KOPF, ["Nirgendwerk", "Phantom", "LL-X"]])), 201)
        alt = bestandsimport.SPERRE_WARTEN
        bestandsimport.SPERRE_WARTEN = 0.1
        assert bestandsimport.SPERRE.acquire(timeout=1)
        try:
            for r in (probelauf(w, t.lager, daten), uebernehmen(w, t.lager, daten),
                      w.post(f"{url}/nacharbeit/pruefen", json={"zeilen": "alle"})):
                assert r.status_code == 429 and "anderer Bestandsimport" in r.json["error"], r.get_json()
            assert w.get(f"{url}/nacharbeit").status_code == 200
        finally:
            bestandsimport.SPERRE.release()
            bestandsimport.SPERRE_WARTEN = alt
        abfragen = []
        get_db = db_mod.get_db

        def mit_zaehler():
            conn = get_db()
            conn.set_trace_callback(abfragen.append)
            return conn
        db_mod.get_db = mit_zaehler
        try:
            t0 = time.time()
            b = ok(probelauf(w, t.lager, daten))
            t1 = time.time()
            n_probe = len(abfragen)
            r = ok(uebernehmen(w, t.lager, daten), 201)
            t2 = time.time()
        finally:
            db_mod.get_db = get_db
        assert b["zahlen"]["uebernehmbar"] == 2000 and r["import"]["angelegt"] == 2000
        print(f"  2000 Zeilen: Probelauf {t1 - t0:.1f} s ({n_probe / 2000:.1f} Abfragen je Zeile), "
              f"Übernehmen {t2 - t1:.1f} s ({(len(abfragen) - n_probe) / 2000:.1f} Abfragen je Zeile)")
        assert n_probe / 2000 < 40 and (len(abfragen) - n_probe) / 2000 < 80
    print("Last und Leistung bestanden.")


# --- 24. Datenschutz und Frist --------------------------------------------------------------------------------------

def test_datenschutz_und_frist():
    from app.material.services import bestandsimport
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        app, w = t.app, t.w
        w2, w2_id = neues_konto(app, "wart2@example.org", name="Wart Zwei")
        freigeben(t.adm, t.gid, w2_id, "verwalten")
        kopf = ["Hersteller", "Modell", "Seriennummer", "Bemerkung", "Ordner", "<b>Titel</b>"]
        r = ok(uebernehmen(w2, t.lager, csv_bytes([kopf, ["Nirgendwerk", "Phantom", "DS-0001", "y" * 6000, "", "x"],
                                                   ["Nirgendwerk", "Phantom <img src=x>", "DS-0002", "",
                                                    "<a href=\"/x\">MTW</a>", "<b>x</b>"]])), 201)
        assert any(s["titel"] == "<b>Titel</b>" for s in r["spalten"])
        assert any("<a href=\"/x\">MTW</a>" in g_["text"] for z in r["zeilen"] for g_ in z["gruende"])
        offen = offene_zeilen(app, t.lager)
        assert len(json.loads(offen[0]["values_json"])["bemerkung"]) == 4100
        ok(w2.post(f"/api/material/ordner/{t.lager}/nacharbeit/verwerfen", json={"zeilen": [offen[0]["id"]]}))
        assert t.adm.delete(f"/api/admin/users/{w2_id}", headers=H).status_code == 200
        assert zeilen(app, "SELECT created_by, created_by_name FROM mat_import")[0] == {"created_by": None,
                                                                                       "created_by_name": "(gelöschtes Konto)"}
        assert zeilen(app, "SELECT resolved_by_name FROM mat_import_row WHERE id = ?",
                      (offen[0]["id"],))[0]["resolved_by_name"] == "(gelöschtes Konto)"
        # Frist (ein Monat): ein Import von vor fünf Wochen verfällt, ein jüngerer bleibt offen.
        ok(uebernehmen(w, t.lager, csv_bytes([KOPF, ["Nirgendwerk", "Phantom", "DS-0003"]])), 201)
        alt = (datetime.now(timezone.utc) - timedelta(days=35)).replace(microsecond=0).isoformat()
        sql(app, "UPDATE mat_import SET created_at = ? WHERE id = (SELECT MIN(id) FROM mat_import)", (alt,))
        conn = sqlite3.connect(os.path.join(tmp, "stroemis.db"))
        assert bestandsimport.aufraeumen(conn) == 1
        conn.close()
        z = zeilen(app, "SELECT * FROM mat_import_row WHERE id = ?", (offen[1]["id"],))[0]
        assert (z["status"], z["values_json"], z["resolved_by_name"]) == ("verworfen", "{}", "System")
        assert "einem Monat" in z["reason"]
        assert len(offene_zeilen(app, t.lager)) == 1
        sql(app, "UPDATE mat_import SET created_at = ?", (alt,))
        with app.app_context():
            assert bestandsimport.aufraeumen() == 1 and bestandsimport.aufraeumen() == 0
        # Ein Ordner mit offener Nacharbeit lässt sich nicht löschen.
        fach = ordner_neu(w, t.lager, "Fach Weg")
        ok(uebernehmen(w, fach, csv_bytes([KOPF, ["Nirgendwerk", "Phantom", "DS-0010"]])), 201)
        r = w.delete(f"/api/material/ordner/{fach}", headers=H)
        assert r.status_code == 409 and "Nacharbeit" in r.json["error"], r.get_json()
        ok(w.post(f"/api/material/ordner/{fach}/nacharbeit/verwerfen", json={"zeilen": "alle"}))
        assert w.delete(f"/api/material/ordner/{fach}", headers=H).status_code == 200
        assert not zeilen(app, "SELECT 1 FROM mat_import WHERE folder_id = ?", (fach,))
    print("Datenschutz und Frist bestanden.")


# --- 25. DEMO -------------------------------------------------------------------------------------------------------

def test_demo():
    with tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        gdemo, _ = gliederung(t.adm, "DEMO Platzhalter")
        sql(t.app, "UPDATE mat_folder SET demo = 1 WHERE id = ?", (gdemo,))
        demo = produkt(t.mv, t.ids["karabiner"], "Demowerk", "Demo Gerät", "DE-1")
        sql(t.app, "UPDATE mat_product SET demo = 1 WHERE id = ?", (demo,))
        r = ok(uebernehmen(t.w, gdemo, csv_bytes([KOPF, ["Demowerk", "Demo Gerät", "DM-0001"],
                                                  ["Testwerk", "HMS Beta", "DM-0002"]])), 201)
        for z in r["zeilen"]:
            assert z["status"] == "angelegt" and t.w.get(f"/api/material/artikel/{z['artikel']['id']}").json["artikel"][
                "demo"], z
        b = ok(probelauf(t.w, t.lager, csv_bytes([KOPF, ["Demowerk", "Demo Gerät", "DM-0003"]])))
        assert b["zeilen"][0]["status"] == "nacharbeit" and b["zeilen"][0]["produkt"] is None
        ok(uebernehmen(t.w, t.lager, csv_bytes([KOPF, ["Demowerk", "Demo Gerät", "DM-0003"]])), 201)
        r = t.w.post(f"/api/material/ordner/{t.lager}/nacharbeit/produkt", json={"zeilen": "alle", "product_id": demo})
        assert r.status_code == 400 and "DEMO" in r.json["error"]
    print("DEMO bestanden.")


# --- 26. Seite und Hinweise -----------------------------------------------------------------------------------------

def test_seite_und_hinweise():
    with ohne_post(), tempfile.TemporaryDirectory() as tmp:
        t = zwei_gliederungen(tmp)
        w, k = t.wart_a, t.k
        r = w.get(f"/material/ordner/{t.lager_a}/import")
        assert r.status_code == 200 and 'data-ansicht="bestandsimport"' in r.get_data(as_text=True)
        assert k["fuehrung"].get(f"/material/ordner/{t.lager_a}/import").status_code == 403
        assert w.get(f"/material/ordner/{t.lager_b}/import").status_code == 404
        assert w.get(f"/material/ordner/{t.sys_a['gesperrt']}/import").status_code == 409
        assert w.get(f"/api/material/ordner/{t.lager_a}").json["ordner"]["nacharbeit"] == 0
        ok(uebernehmen(w, t.lager_a, csv_bytes([KOPF, ["Nirgendwerk", "Phantom", "SH-0001"]])), 201)
        ok(uebernehmen(t.wart_b, t.lager_b, csv_bytes([KOPF, ["Nirgendwerk", "Phantom", "SH-0002"]])), 201)
        assert w.get(f"/api/material/ordner/{t.lager_a}").json["ordner"]["nacharbeit"] == 1
        assert k["fuehrung"].get(f"/api/material/ordner/{t.lager_a}").json["ordner"]["nacharbeit"] is None
        start = w.get("/api/material/start").json["nacharbeit"]
        assert start == [{"ordner": {"id": t.lager_a, "pfad": "OG Musterstadt / Lager Musterstadt"}, "offen": 1}], start
        assert k["fuehrung"].get("/api/material/start").json["nacharbeit"] == []
        assert len(k["katalog"].get("/api/material/start").json["nacharbeit"]) == 2
    print("Seite und Hinweise bestanden.")


# --- 27. Befunde der Prüfung vor Auslieferung (Importlogik) ----------------------------------------------------------

def test_befunde_vor_auslieferung():
    """Die Befunde der Prüfung „Importlogik“ vor der Auslieferung (CHANGELOG „Prüfung vor Auslieferung“): Seriennummern-
    Synonyme und Warnung ohne Seriennummernspalte, Einheit des Prüfintervalls, dieselbe Liste noch einmal bei offener
    Nacharbeit, Grundprodukt mit Varianten ohne Artikelnummer, gleiche Seriennummer bei zwei Varianten in einer Datei,
    Nacharbeits-CSV nach Ablauf der Frist, übliche Schreibweisen der Anzahl, Lebensende bei Metall; dazu die Sperre
    derselben Datei im Probelauf, die Zahl zeilenweise zugeordneter Zeilen je Karte und die Seriennummer erledigter
    Zeilen."""
    from app.material.services import bestandsimport as bi
    # Lesen des Intervalls: Einheit aus dem Wert, Wörter; 1 und 2 ohne Einheit sind unklar.
    for text, soll in (("12", (12, None)), ("12 Monate", (12, None)), ("1 Monat", (1, None)), ("12M", (12, None)),
                       ("1 Jahr", (12, None)), ("2 Jahre", (24, None)), ("0,5 Jahre", (6, None)), ("1 J.", (12, None)),
                       ("jährlich", (12, None)), ("1x jährlich", (12, None)), ("2 x jährlich", (6, None)),
                       ("halbjährlich", (6, None)), ("vierteljährlich", (3, None)), ("2-jährlich", (24, None)),
                       ("alle 12 Monate", (12, None)), ("alle 2 Jahre", (24, None)), ("1", (None, "einheit")),
                       ("2", (None, "einheit")), ("3", (3, None)), ("ca. jährlich", (None, "unlesbar")),
                       ("0", (None, "unlesbar")), ("11 Jahre", (None, "unlesbar"))):
        assert bi._intervall_lesen(text) == soll, (text, bi._intervall_lesen(text))
    for titel, soll in (("Prüfintervall (Jahre)", "jahre"), ("Intervall in Jahren", "jahre"), ("Prüfintervall [J]", "jahre"),
                        ("Prüfintervall (Monate)", "monate"), ("Prüfzyklus [M]", "monate"), ("Prüfintervall", None)):
        assert bi.intervall_einheit(titel) == soll, titel
    for text, soll in (("1", 1), ("01", 1), ("1,0", 1), ("1 Stk.", 1), ("1 Stück", 1), ("2 Stk.", 2), ("ca. 5", None)):
        assert bi._anzahl_lesen(text) == soll, text
    with ohne_post(), tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        app, w, kar = t.app, t.w, t.ids["karabiner"]
        # 1. Seriennummer: Synonyme; ohne zugeordnete Spalte eine deutliche Warnung mit den nicht zugeordneten Spalten.
        for titel in ("Serien-Nummer", "Ser.-Nr.", "Ser. Nr.", "Fabr.-Nr."):
            assert bi.erkennen(titel)[0] == "seriennummer", titel
        r = ok(uebernehmen(w, t.lager, csv_bytes([["Hersteller", "Modell", "Serien-Nummer"],
                                                  ["Testwerk", "HMS Beta", "BV-0001"]])), 201)
        assert r["zeilen"][0]["seriennummer"] == "BV-0001" and not r["keine_seriennummer"]
        b = ok(probelauf(w, t.lager, csv_bytes([["Hersteller", "Modell", "Kennung", "Ausgegeben an"],
                                                ["Testwerk", "HMS Beta", "BV-0002", "Platzhalter Person"]])))
        assert b["keine_seriennummer"] and b["warnungen"][0].startswith(
            "Keine Spalte ist als Seriennummer zugeordnet – alle 1 Zeilen bekämen eine Seriennummer vom System"), b
        assert "Nicht zugeordnet: „Kennung“." in b["warnungen"][0] and "Ausgegeben" not in b["warnungen"][0]
        # 2. Prüfintervall: Einheit aus dem Spaltentitel; 1 ohne Einheit in die Nacharbeit; kürzer übernommen mit Hinweis.
        b = ok(probelauf(w, t.lager, csv_bytes([["Hersteller", "Modell", "Seriennummer", "Prüfintervall (Jahre)"],
                                                ["Testwerk", "HMS Beta", "BV-0010", "1"],
                                                ["Testwerk", "HMS Beta", "BV-0011", "0,5"]])))
        z = nach_nr(b)
        assert z[2]["status"] == "uebernehmbar" and any("nicht kürzer" in h for h in z[2]["hinweise"]), z[2]
        assert z[3]["status"] == "uebernehmbar" and any("Intervall laut Liste (6 Monate) kürzer" in h
                                                         for h in z[3]["hinweise"]), z[3]
        r = ok(uebernehmen(w, t.lager, csv_bytes([["Hersteller", "Modell", "Seriennummer", "Prüfintervall"],
                                                  ["Testwerk", "HMS Beta", "BV-0012", "1"],
                                                  ["Testwerk", "HMS Beta", "BV-0013", "halbjährlich"],
                                                  ["Testwerk", "HMS Beta", "BV-0014", "jährlich"]])), 201)
        z = nach_nr(r)
        assert codes(z[2]) == ["intervall"] and "Monate oder Jahre?" in z[2]["gruende"][0]["text"]
        assert z[3]["status"] == z[4]["status"] == "angelegt"
        iv = {i["serial_no"]: i["interval_months"] for i in zeilen(app, "SELECT serial_no, interval_months FROM mat_item")}
        assert iv["BV-0013"] == 6 and iv["BV-0014"] is None, iv
        # Die Einheit bleibt in der Nacharbeit (und ihrer CSV) erhalten: „1“ unter „Prüfintervall (Jahre)“ ist „1 Jahr“.
        ok(uebernehmen(w, t.lager, csv_bytes([["Hersteller", "Modell", "Seriennummer", "Prüfintervall (Jahre)"],
                                              ["Nirgendwerk", "Phantom", "BV-0015", "1"]])), 201)
        assert json.loads(offene_zeilen(app, t.lager)[-1]["values_json"])["intervall"] == "1 Jahr"
        ok(w.post(f"/api/material/ordner/{t.lager}/nacharbeit/verwerfen", json={"zeilen": "alle"}))
        # 3. Dieselbe Liste noch einmal, solange ihre Zeilen in der Nacharbeit stehen: ersetzt, nicht verdoppelt – auch
        # mit der Produktwahl aus der Nacharbeit. Binnen 10 Minuten nennt der Probelauf, ab wann es wieder geht.
        hiviz_m = produkt(t.mv, kar, "Platzwerk", "Rahmen | M", "")
        produkt(t.mv, kar, "Platzwerk", "Rahmen | L", "")
        liste = csv_bytes([["Hersteller", "Modell", "Seriennummer", "Bemerkung"],
                           ["Nirgendwerk", "Phantom W", "", "Platzhalter eins"], ["Nirgendwerk", "Phantom W", "", ""],
                           ["Nirgendwerk", "Phantom W", "DN-0003", ""], ["Platzwerk", "Rahmen", "DN-0004", "Größe M"]])
        r = ok(uebernehmen(w, t.lager, liste), 201)
        assert r["import"]["nacharbeit"] == 4 and r["import"]["angelegt"] == 0
        b = ok(probelauf(w, t.lager, liste))
        assert b["sperre_bis"] and "Noch einmal übernehmen geht erst ab" in b["warnungen"][0], b["warnungen"]
        produkt(t.mv, kar, "Nirgendwerk", "Phantom W", "")
        rahmen = [o for o in offene_zeilen(app, t.lager) if o["line_no"] == 5][0]
        ok(w.post(f"/api/material/ordner/{t.lager}/nacharbeit/produkt", json={"zeilen": [rahmen["id"]],
                                                                            "product_id": hiviz_m}))
        sql(app, "UPDATE mat_import SET created_at = ?",
            ((datetime.now(timezone.utc) - timedelta(days=1)).replace(microsecond=0).isoformat(),))
        vorher = len(zeilen(app, "SELECT 1 FROM mat_item"))
        b = ok(probelauf(w, t.lager, liste))
        assert b["sperre_bis"] is None and "4 Zeilen dieser Datei stehen noch in der Nacharbeit" in b["warnungen"][0]
        assert "noch einmal angelegt" not in b["warnungen"][0] and b["zahlen"]["uebernehmbar"] == 4, b["warnungen"]
        r = ok(uebernehmen(w, t.lager, liste), 201)
        assert (r["import"]["angelegt"], r["import"]["ersetzt"]) == (4, 4) and not offene_zeilen(app, t.lager)
        assert len(zeilen(app, "SELECT 1 FROM mat_item")) == vorher + 4
        assert nach_nr(r)[5]["produkt"]["id"] == hiviz_m
        # 4. Grundprodukt und Variante ohne Artikelnummer: „Vertex Platz“ ist mehrdeutig, die Variante selbst genau.
        produkt(t.mv, kar, "Kopfwerk", "Vertex Platz", "")
        hiviz = produkt(t.mv, kar, "Kopfwerk", "Vertex Platz | Hi-Viz", "")
        b = ok(probelauf(w, t.lager, csv_bytes([KOPF, ["Kopfwerk", "Vertex Platz", "BV-0020"],
                                                ["Kopfwerk", "Vertex Platz | Hi-Viz", "BV-0021"],
                                                ["Kopfwerk", "Vertex Platz Hi-Viz", "BV-0022"]])))
        z = nach_nr(b)
        assert codes(z[2]) == ["produkt_mehrdeutig"] and z[3]["produkt"]["id"] == z[4]["produkt"]["id"] == hiviz
        # 5. Dieselbe Seriennummer bei zwei Varianten desselben Herstellers in einer Datei – und in der Auswahl.
        produkt(t.mv, kar, "Baumwerk", "Sequ", "SQ-1")
        produkt(t.mv, kar, "Baumwerk", "Sequ | Gross", "SQ-2")
        r = ok(uebernehmen(w, t.lager, csv_bytes([["Hersteller", "Modell", "Artikelnummer", "Seriennummer"],
                                                  ["Baumwerk", "Sequ", "SQ-1", "VAR-0002"],
                                                  ["Baumwerk", "Sequ | Gross", "SQ-2", "VAR-0002"]])), 201)
        z = nach_nr(r)
        assert codes(z[2]) == codes(z[3]) == ["serie_variante"] and "(Zeilen 3)" in z[2]["gruende"][0]["text"], z[2]
        assert not zeilen(app, "SELECT 1 FROM mat_item WHERE serial_no = 'VAR-0002'")
        e = ok(w.post(f"/api/material/ordner/{t.lager}/nacharbeit/pruefen", json={"zeilen": "alle"}))
        assert [codes(x) for x in e["zeilen"]] == [["serie_variante"]] * 2
        assert "in der Auswahl" in e["zeilen"][0]["gruende"][0]["text"], e["zeilen"][0]
        ok(w.post(f"/api/material/ordner/{t.lager}/nacharbeit/verwerfen", json={"zeilen": "alle"}))
        # 6. Nacharbeits-CSV nach Ablauf der Frist: Die Zeilen werden neu geprüft; von Hand verworfene bleiben erledigt
        # und sagen, wie es doch geht.
        ok(uebernehmen(w, t.lager, csv_bytes([KOPF, ["Nirgendwerk", "Phantom Z", "BV-0030"]])), 201)
        ok(uebernehmen(w, t.lager, csv_bytes([KOPF, ["Nirgendwerk", "Phantom Y", "BV-0031"]])), 201)
        csv_alt = w.get(f"/material/ordner/{t.lager}/nacharbeit.csv").data
        frist, hand = offene_zeilen(app, t.lager)
        ok(w.post(f"/api/material/ordner/{t.lager}/nacharbeit/verwerfen", json={"zeilen": [hand["id"]]}))
        sql(app, "UPDATE mat_import SET created_at = ? WHERE id = ?",
            ((datetime.now(timezone.utc) - timedelta(days=40)).replace(microsecond=0).isoformat(), frist["import_id"]))
        with app.app_context():
            assert bi.aufraeumen() == 1
        produkt(t.mv, kar, "Nirgendwerk", "Phantom Z", "")
        produkt(t.mv, kar, "Nirgendwerk", "Phantom Y", "")
        r = ok(uebernehmen(w, t.lager, csv_alt, "nach_frist.csv"), 201)
        z = {x["seriennummer"]: x for x in r["zeilen"]}
        assert z["BV-0030"]["status"] == "angelegt" and "nach Ablauf der Frist verworfen" in z["BV-0030"]["hinweise"][0]
        assert z["BV-0031"]["status"] == "erledigt" and "Nacharbeit-Nr. leeren" in z["BV-0031"]["hinweise"][0]
        # 7. Anzahl „01“, „1,0“, „1 Stk.“ sind eins; 8. ein Lebensende bei Metall steht in der Bemerkung.
        r = ok(uebernehmen(w, t.lager, csv_bytes([KOPF + ["Anzahl", "Lebensende"],
                                                  ["Testwerk", "HMS Beta", "BV-0040", "01", ""],
                                                  ["Testwerk", "HMS Beta", "BV-0041", "1,0", ""],
                                                  ["Testwerk", "HMS Beta", "BV-0042", "1 Stk.", "2030-01-01"],
                                                  ["Testwerk", "HMS Beta", "BV-0043", "2 Stk.", ""]])), 201)
        z = nach_nr(r)
        assert [z[n]["status"] for n in range(2, 6)] == ["angelegt"] * 3 + ["nacharbeit"]
        assert "Anzahl 2" in z[5]["gruende"][0]["text"] and any("steht in der Bemerkung" in h for h in z[4]["hinweise"])
        assert zeilen(app, "SELECT notes FROM mat_item WHERE serial_no = 'BV-0042'")[0]["notes"] == \
            "Lebensende laut Liste: 2030-01-01"
        # Karte: zeilenweise zugeordnet; erledigte Zeilen tragen ihre Seriennummer.
        ok(w.post(f"/api/material/ordner/{t.lager}/nacharbeit/verwerfen", json={"zeilen": "alle"}))
        ok(uebernehmen(w, t.lager, csv_bytes([KOPF, ["Platzwerk", "Rahmen", "BV-0050"], ["Platzwerk", "Rahmen", "BV-0051"]])),
           201)
        offen = offene_zeilen(app, t.lager)
        ok(w.post(f"/api/material/ordner/{t.lager}/nacharbeit/produkt", json={"zeilen": [offen[0]["id"]],
                                                                            "product_id": hiviz_m}))
        karte = ok(w.get(f"/api/material/ordner/{t.lager}/nacharbeit"))["gruppen"][0]
        assert (karte["offen"], karte["zugeordnet"], karte["produkt"]) == (2, 1, None), karte
        csv_neu = w.get(f"/material/ordner/{t.lager}/nacharbeit.csv").data
        ok(w.post(f"/api/material/ordner/{t.lager}/nacharbeit/verwerfen", json={"zeilen": "alle"}))
        b = ok(probelauf(w, t.lager, csv_neu, "erledigt.csv"))
        assert [(x["status"], x["seriennummer"]) for x in b["zeilen"]] == [("erledigt", "BV-0050"),
                                                                           ("erledigt", "BV-0051")], b["zeilen"]
    print("Befunde vor Auslieferung bestanden.")


# --- 28. Oberfläche (statisch) -------------------------------------------------------------------------------------

# --- Entscheidung vom 20.10.2026: ohne Seriennummernspalte nur mit Bestätigung ---------------------------------------

def test_ohne_seriennummernspalte():
    """Entscheidung vom 20.10.2026: Ist keine Spalte als Seriennummer zugeordnet, verlangt Übernehmen das Feld
    „ohne_seriennummern“ – ohne es 409 und nichts geschrieben, auch wenn alle Zeilen in die Nacharbeit gingen (dort
    bekämen sie ihre STR-Nummer später ohne Frage). Der Probelauf sagt es mit „ohne_seriennummern_bestaetigen“. Einzelne
    Zeilen ohne Nummer bei vorhandener Spalte brauchen nichts. Die Nacharbeits-CSV folgt derselben Regel; „erneut prüfen
    und übernehmen“ fragt nicht noch einmal – bestätigt wurde beim Import."""
    from app.material.services import bestandsimport as bi
    with ohne_post(), tempfile.TemporaryDirectory() as tmp:
        t = aufbau(tmp)
        app, w, hms = t.app, t.w, t.ids["hms"]
        kopf = ["Hersteller", "Modell", "Kennung"]
        ohne = csv_bytes([kopf, ["Testwerk", "HMS Beta", "OS-0001"], ["Testwerk", "HMS Beta", "OS-0002"],
                          ["Nirgendwerk", "Phantom", "OS-0003"]])
        b = ok(probelauf(w, t.lager, ohne))
        assert b["keine_seriennummer"] and b["ohne_seriennummern_bestaetigen"], b
        assert (b["zahlen"]["uebernehmbar"], b["zahlen"]["nacharbeit"], b["zahlen"]["ohne_serie"]) == (2, 1, 2)
        assert "Übernehmen nur mit ausdrücklicher Bestätigung" in b["warnungen"][0], b["warnungen"]
        # Ohne das Feld, mit „false“ oder nur dem alten „ohne_serie“: 409 mit Klartext, nichts geschrieben.
        vorher = stand_aller_tabellen(app)
        for felder in ({}, {"ohne_seriennummern": "false"}, {"ohne_seriennummern": ""}, {"ohne_serie": "1"}):
            r = uebernehmen(w, t.lager, ohne, **felder)
            assert r.status_code == 409 and r.json["error"] == bi.OHNE_SERIENNUMMERN, (felder, r.get_json())
        r = uebernehmen(w, t.lager, ohne, ohne_seriennummern="vielleicht")
        assert r.status_code == 400 and "ja oder nein" in r.json["error"], r.get_json()
        assert stand_aller_tabellen(app) == vorher
        # Mit dem Feld: 201, die Artikel bekommen STR-Nummern, die Zeile ohne Produkt geht in die Nacharbeit.
        r = ok(uebernehmen(w, t.lager, ohne, ohne_seriennummern="true"), 201)
        assert (r["import"]["angelegt"], r["import"]["nacharbeit"]) == (2, 1) and not r["ohne_seriennummern_bestaetigen"]
        assert [z["seriennummer"][:4] for z in r["zeilen"] if z["status"] == "angelegt"] == ["STR-", "STR-"]
        assert not any("Keine Spalte" in x for x in r["warnungen"]) and "haben eine Seriennummer" in r["warnungen"][-1]
        # Die Nacharbeit weiß nichts mehr von der Spalte – „erneut prüfen und übernehmen“ braucht kein Feld.
        offen = offene_zeilen(app, t.lager)
        assert len(offen) == 1 and "seriennummer" not in json.loads(offen[0]["values_json"])
        ok(w.post(f"/api/material/ordner/{t.lager}/nacharbeit/produkt", json={"zeilen": [offen[0]["id"]],
                                                                                "product_id": hms}))
        e = ok(w.post(f"/api/material/ordner/{t.lager}/nacharbeit/pruefen", json={"zeilen": "alle", "uebernehmen": True}))
        assert e["zahlen"]["angelegt"] == 1 and e["zeilen"][0]["seriennummer"].startswith("STR-"), e
        # Nur Nacharbeit (kein Artikel übernehmbar): ebenso nur mit Bestätigung.
        nur_na = csv_bytes([kopf, ["Nirgendwerk", "Phantom", "OS-0010"]])
        b = ok(probelauf(w, t.lager, nur_na, "nur_nacharbeit.csv"))
        assert (b["zahlen"]["uebernehmbar"], b["zahlen"]["nacharbeit"]) == (0, 1) and b["ohne_seriennummern_bestaetigen"]
        assert uebernehmen(w, t.lager, nur_na, "nur_nacharbeit.csv").status_code == 409
        r = ok(uebernehmen(w, t.lager, nur_na, "nur_nacharbeit.csv", ohne_seriennummern="1"), 201)
        assert r["import"]["nacharbeit"] == 1
        # Spalte vorhanden, einzelne Zeilen leer: wie entschieden ohne Bestätigung (STR für die leere Zeile).
        mit = csv_bytes([KOPF, ["Testwerk", "HMS Beta", "OS-0020"], ["Testwerk", "HMS Beta", ""]])
        b = ok(probelauf(w, t.lager, mit, "mit.csv"))
        assert not b["keine_seriennummer"] and not b["ohne_seriennummern_bestaetigen"] and b["zahlen"]["ohne_serie"] == 1
        r = ok(uebernehmen(w, t.lager, mit, "mit.csv"), 201)
        assert r["import"]["angelegt"] == 2 and r["zeilen"][1]["seriennummer"].startswith("STR-")
        # Die Spalte von Hand als Seriennummer zugeordnet: keine Bestätigung, die Nummer der Liste gilt.
        kennung = csv_bytes([kopf, ["Testwerk", "HMS Beta", "OS-0030"]])
        b = ok(probelauf(w, t.lager, kennung, "kennung.csv", zuordnung={"2": "seriennummer"}))
        assert not b["ohne_seriennummern_bestaetigen"]
        r = ok(uebernehmen(w, t.lager, kennung, "kennung.csv", zuordnung={"2": "seriennummer"}), 201)
        assert r["zeilen"][0]["seriennummer"] == "OS-0030"
        # Nacharbeits-CSV: mit ihrer Spalte „Seriennummer“ (hier leer) keine Bestätigung; ohne die Spalte – gelöscht
        # oder auf „ignorieren“ gestellt – dieselbe Regel wie jede Liste.
        tab = _csv_lesen(w.get(f"/material/ordner/{t.lager}/nacharbeit.csv").data)
        assert tab[0][5] == "Seriennummer" and len(tab) == 2 and tab[1][5] == "", tab
        b = ok(probelauf(w, t.lager, csv_bytes(tab), "nacharbeit.csv"))
        assert b["nacharbeit_csv"] and not b["keine_seriennummer"] and not b["ohne_seriennummern_bestaetigen"]
        b = ok(probelauf(w, t.lager, csv_bytes(tab), "nacharbeit.csv", zuordnung={"5": "ignorieren"}))
        assert b["nacharbeit_csv"] and b["ohne_seriennummern_bestaetigen"]
        ohne_spalte = csv_bytes([z[:5] + z[6:] for z in tab])
        b = ok(probelauf(w, t.lager, ohne_spalte, "nacharbeit.csv"))
        assert b["nacharbeit_csv"] and b["ohne_seriennummern_bestaetigen"] and b["zahlen"]["nacharbeit"] == 1
        r = uebernehmen(w, t.lager, ohne_spalte, "nacharbeit.csv")
        assert r.status_code == 409 and r.json["error"] == bi.OHNE_SERIENNUMMERN
        assert len(offene_zeilen(app, t.lager)) == 1
        r = ok(uebernehmen(w, t.lager, ohne_spalte, "nacharbeit.csv", ohne_seriennummern="true"), 201)
        assert (r["import"]["ersetzt"], r["import"]["nacharbeit"]) == (1, 1)
    print("Ohne Seriennummernspalte bestanden.")


def test_oberflaeche_statisch():
    """material.js (Paket OBERFLÄCHE): Ansicht „bestandsimport“ im Reiter „Bestand“ mit Ordnerleiste; Einstieg über das
    Untermenü „Export / Import“ im Ordnerkopf, Hinweise im Ordnerkopf und auf der Übersicht; ein roter Knopf
    „Übernehmen: …“ mit Zahlen, Probelauf abbrechbar (AbortController), 413/429 im Klartext; Nacharbeit mit Filter nach
    Grund, „alle n offenen“ (→ "alle"), Frist und offen_bis; Melden über die Nacharbeit (angabe, ohne folder_id),
    „Im Katalog anlegen …“ ohne Seitenwechsel; Knopfregel (kein btn ghost, danger nur über die Rückfrage)."""
    wurzel = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    js = open(os.path.join(wurzel, "app/material/static/material.js"), encoding="utf-8").read()
    css = open(os.path.join(wurzel, "app/material/static/material.css"), encoding="utf-8").read()
    assert 'bestandsimport: "bestand" };' in js and '"meldungen", "bestandsimport"]);' in js
    assert re.search(r"const ANSICHTEN = \{[^}]*\bbestandsimport \};", js), "Ansicht fehlt in ANSICHTEN"
    anfang = js.index("  async function bestandsimport() {")
    block = js[anfang:js.index("\n  }\n", anfang)]
    # Einstieg und Hinweise
    assert '${importieren ? "Export / Import" : "Export"}' in js and "exportMenue(d, darfPflegen)" in js
    assert "<strong>Bestandsliste importieren …</strong>" in js and "<strong>Nacharbeit (${d.nacharbeit} offen)</strong>" in js
    assert "auf Nacharbeit." in js and "Nacharbeit öffnen</a>" in js and "${nacharbeitHinweis(d.nacharbeit)}" in js
    # Probelauf, Übernehmen, Fehler
    assert "new AbortController()" in block and "lauf.abort()" in block and "setTimeout(probelauf, 600)" in block
    assert 'return `Übernehmen: ${a} anlegen, ${n} in die Nacharbeit`;' in block
    assert 'return `${n} in die Nacharbeit übernehmen`;' in block
    assert "fehler.status === 429" in block and 'data-act="bi-wiederholen">Erneut versuchen' in block
    assert "e.status === 413 ? IMPORT_ZU_GROSS" in js and "f.size > IMPORT_MAX_BYTES" in block
    assert "IMPORT_MAX_BYTES = 5 * 1024 * 1024" in js
    assert "angelegtTabelle(teil)" in block and "IMPORT_NEUBESCHRIFTEN" in block
    # Nacharbeit
    assert 'gruppe: naGruppe || "alle"' in block and '["produkt", "Produkt"], ["seriennummer", "Seriennummer"]' in js
    assert 'naAlle ? "alle" : Array.from(naWahl)' in block and "esc(fmtTag(z.offen_bis))" in block
    assert "esc(na.frist)" in block and "esc(b.frist)" in block
    assert "url: `${API}/ordner/${fid}/nacharbeit/melden`" in block and "angabe: g.angabe" in block
    assert "if (v.url) felder.angabe = v.angabe" in js and "await api(v.url || `${API}/meldungen`" in js
    assert "if (opt.danach) await opt.danach(r.produkt);" in js
    assert "/nacharbeit/produkt`, { method: \"POST\", body: { angabe: g.angabe, product_id: p.id } }" in block
    assert "if (!g.produkt && DARF.katalog) knoepfe.push(" in block and "Im Katalog anlegen …" in block
    # Knopfregel: rot nur „Übernehmen: …“ (Seite), „Zuordnen“ und „n Artikel übernehmen“ (je ein Dialog)
    assert block.count('class="btn"') == 2 and 'class="btn" data-act="bi-uebernehmen"' in block, block.count('class="btn"')
    assert block.count('class="btn mat-knopf-umbruch"') == 1 and "btn danger" not in block
    assert "btn ghost" not in js
    # Dateiinhalte nur maskiert: Gründe, Hinweise, Beispiele
    assert '<li data-code="${esc(x.code)}">${esc(x.text)}</li>' in js and "<li>${esc(x)}</li>" in js
    assert '<span class="mat-beispiel">${esc(w)}</span>' in block
    # Schmale Ansichten über den Container „mat-seite“; der lange Knopf bricht um
    teil = css[css.index("/* Phase 2 · Lieferung 3"):]
    assert "@container mat-seite (max-width: 600px)" in teil and ".mat-import-knoepfe .btn { width: 100%; }" in teil
    assert ".mat-import-knoepfe .btn, .mat-knopf-umbruch { white-space: normal;" in teil
    # Prüfung vor Auslieferung (Browser): 1 Ergebnis „Übernommen“ im Dialog breit und mit eigenem Behälter (Karten am
    # Telefon); 2 dieselbe Datei gesperrt – kein „Übernehmen“, sondern ab wann; 3 Nummern und Daten am Stück; 4 Fokus
    # nach dem Neuzeichnen; 5/6 Karten mit zeilenweiser Zuordnung bzw. schon zugeordnet; 7 „–“ statt „ohne“ bei
    # erledigten Zeilen; 8 Bemerkung mit Zeilenumbrüchen; 9 Telefon: „Datei wählen“ breit, Kästchen links oben.
    assert '<div class="mat-dlg-liste">${angelegtTabelle(angelegt, true)}</div>' in block and '"mat-uebernommen"' in block
    assert ".mat-dlg-liste { container: mat-seite / inline-size; }" in teil and "dialog:where(.mat-uebernommen)" in teil
    assert 'const text = fehlt || sperre ? "" : knopfText(b);' in block and "noch einmal übernehmen geht erst ab" in block
    assert "|| (bericht && bericht.sperre_bis)) return;" in block
    assert "table.list:is(.mat-import-zeilen, .mat-angelegt) .mat-ungeteilt { white-space: nowrap; overflow-wrap: normal; }" in teil
    assert js.count('mat-ungeteilt') >= 5
    assert "if (titel) { titel.tabIndex = -1; titel.focus(); }" in block and '$("#nz-ok", body).focus();' in block
    assert 'const fokus = fertig ? $(".mat-austausch-stand", berichtEl)' in block and 'ok" tabindex="-1">Übernommen:' in block
    assert "if (neu) neu.focus();" in js[js.index("  function wikiWahl(formular, danach) {"):js.index("/* --- Suche ---")]
    assert "zeilenweise zugeordnet: ${g.zugeordnet} von ${g.offen}" in block and "Das ersetzt die zeilenweise Zuordnung" in block
    assert 'esc(g.produkt ? "zugeordnet" : PRODUKT_ART[g.art] || g.art)' in block and "if (!g.produkt && DARF.katalog)" in block
    assert '${x.status === "erledigt" ? "–" : "ohne"}' in block and "white-space: pre-line;" in teil
    # Entscheidung vom 20.10.2026: ohne Seriennummernspalte ein Kästchen direkt über dem roten Knopf; der Knopf bleibt
    # gesperrt, bis es gesetzt ist, und erst dann geht „ohne_seriennummern“ mit. Eine neue Datei, jeder neue Probelauf,
    # eine geänderte Zuordnung und „Neu wählen“ nehmen das Häkchen zurück; es setzt den Knopf frei, ohne neu zu zeichnen.
    knoepfe = block[block.index("${text && b.ohne_seriennummern_bestaetigen ?"):]
    assert (knoepfe.index('id="bi-ohne-serie"') < knoepfe.index('<div class="btn-row mat-import-knoepfe">')
            < knoepfe.index('data-act="bi-uebernehmen"'))
    assert ('<input type="checkbox" id="bi-ohne-serie" aria-describedby="bi-ohne-serie-hilfe" ${ohneSerie ? "checked" : ""}>'
            in block) and 'id="bi-ohne-serie-hilfe">${esc(IMPORT_OHNE_SERIENSPALTE)}' in block
    assert "|| Boolean(bericht && bericht.ohne_seriennummern_bestaetigen && !ohneSerie);" in block
    assert 'data-act="bi-uebernehmen" ${uebernehmenGesperrt() ? "disabled" : ""}' in block
    assert "if (!datei || uebernehmenGesperrt() || (bericht && bericht.sperre_bis)) return;" in block
    assert 'if (bericht && bericht.ohne_seriennummern_bestaetigen) fd.append("ohne_seriennummern", "true");' in block
    assert "ohneSerie = false;" in block[block.index("async function probelauf() {"):block.index("function dateiZeile(b) {")]
    assert block.count("ohneSerie = false;") == 5 and 'if (h) h.checked = false;' in block   # 1 + 4 Rücksetzer
    assert "ohneSerie = ev.target.checked;\n        if (k) k.disabled = uebernehmenGesperrt();\n        return;" in block
    assert "keine_seriennummer" not in block
    assert "`Diese Liste hat keine Seriennummern – ${artikel} eine Seriennummer vom System (STR-…).`" in js
    assert ".mat-import-ohne-serie label.inline { display: flex; align-items: flex-start;" in teil
    assert ".mat-import .mat-ablage-knopf { display: flex;" in teil and "table.list.mat-na-zeilen td.mat-k-wahl { order: -3;" in teil
    print("Oberfläche (statisch) bestanden.")


if __name__ == "__main__":
    test_migration_12()
    test_vereinsliste_ende_zu_ende()
    test_lesen_csv()
    test_lesen_xlsx_und_grenzen()
    test_seriennummer()
    test_spalten()
    test_daten_und_blockieren()
    test_faelligkeit()
    test_produktabgleich()
    test_probelauf_schreibt_nichts()
    test_uebernehmen()
    test_wiederholung()
    test_nacharbeit_csv_rundweg()
    test_varianten_zeilenweise()
    test_sammelzeilen()
    test_ordnerspalte()
    test_pflichtangaben_und_regeln()
    test_rueckruf_im_probelauf()
    test_fremde_dublette_und_abzaehlen()
    test_rechte()
    test_nacharbeit()
    test_produkt_melden()
    test_last_und_leistung()
    test_datenschutz_und_frist()
    test_demo()
    test_seite_und_hinweise()
    test_befunde_vor_auslieferung()
    test_ohne_seriennummernspalte()
    test_oberflaeche_statisch()
    print("Alle Tests zum Bestandsimport bestanden.")

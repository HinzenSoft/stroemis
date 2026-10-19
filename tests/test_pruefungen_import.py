"""Prüft das Excel-Import-Modul ohne Flask und ohne Datenbank:
python tests/test_pruefungen_import.py  oder  pytest tests/test_pruefungen_import.py

Die Excel-Dateien entstehen hier im Test mit openpyxl – strukturgleich zum Export des
Seminarsystems (Kopfpaare mit verbundenen Zellen, Kopfzeile in Zeile 7, mehrzeilige Titel),
aber ausschließlich mit erfundenen Namen."""
import io
import os
import sys
import time
import tracemalloc
import zipfile
from datetime import date, datetime

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import pruefungen_import as pi  # noqa: E402
from app.pruefungen_import import (  # noqa: E402
    ImportFehler, ZUORDNUNGEN, analysieren, datum_lesen, ist_erfuellt, ist_leitungsrolle, normalisiert, passt,
    schluessel, zeitraum_lesen,
)

VERWALTUNG = ["Bestätigung TN", "Bestätigung Gliederung", "Kostenübernahme", "Zahlung"]
VORAUSSETZUNGEN = [
    "Mindestalter 16 Jahre zu Beginn der Veranstaltung",
    "Mitgliedschaft in der DLRG\n(wird automatisch geprüft bzw. bei der Heimatgliederung angefragt)",
    "Ärztliche Tauglichkeitsuntersuchung oder Selbsterklärung zur Gesundheit",
    "Erfolgreiche Teilnahme am Coopertest (Prüfung zu Lehrgangsbeginn)",
    "Deutsches Rettungsschwimmabzeichen Silber (152), nicht älter als 2 Jahre zu Beginn",
    "Sanitätsausbildung A (331), nicht älter als 4 Jahre zu Beginn",
    "Basisausbildung Einsatzdienste (401) vorhanden zum Meldeschluss",
    "Modul - Umgang mit Rettungsgeräten (402) vorhanden zum Meldeschluss",
    "Modul - Schwimmen in fließenden Gewässern (403) vorhanden zum Meldeschluss",
    "400m schwimmen in 8 Minuten\n(Nachweis kann vor Ort erbracht werden)",
    "Einverständniserklärung bei minderjährigen Teilnehmenden",
    "Für RBK: Erfolgreiche Teilnahme am Sichtungstreffen\nFür Externe: Nachweis der Gliederung",
    "Abschlusstest Theoriemodul (einzureichen bis 28.11.2026, 23:59 Uhr)",
    "Selbsterklärung zur persönlichen Schutzausrüstung",
    "Beherrschen der Knoten aus der 401-Basisausbildung Einsatzdienste",
]
KOPF = ["Vorname", "Nachname", "Alter", "Rolle", "Gliederung", "Status"] + VERWALTUNG + VORAUSSETZUNGEN

# Erfundene Teilnehmende: (Vorname, Nachname, Geburtsdatum als Zelltext bzw. datetime, Rolle, Gliederung)
PERSONEN = [
    ("Anna", "Beispiel", "2002-10-03 00:00:00", "Teilnehmender", "Musterstadt"),
    ("Bernd", "Probe", "1999-01-15 00:00:00", "Teilnehmender", "Musterstadt"),
    ("Clara", "Test", datetime(2005, 6, 30), "Teilnehmender", "Beispielhausen"),
    ("Dirk", "Muster", "1988-12-24 00:00:00", "Lehrgangsleitung", "Musterstadt"),
    ("Eva", "Fiktiv", "2001-02-28 00:00:00", "Teilnehmender", "Beispielhausen"),
    ("Anna", "Beispiel", "2002-10-03 00:00:00", "Teilnehmender", "Musterstadt"),   # Duplikat von Zeile 8
    ("", "Frank Erfunden", "2000-05-05 00:00:00", "Teilnehmender", "Musterstadt"),  # voller Name in „Nachname“
]


def beispiel_xlsx(personen=PERSONEN, kopf=KOPF, kopfzeile=7, leerzeilen=2, kopfpaare=None):
    """Baut die Datei im Speicher – gleiche Form wie der echte Export, nur mit erfundenen Daten."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Worksheet"
    if kopfpaare is None:
        kopfpaare = [("Nr.", "2026-0008"), ("Titel", "Strömungsretter 1 (SR1)"),
                     ("Zeitraum", "04.12.2026 18:00 Uhr bis 13.12.2026 19:00 Uhr"),
                     ("Ort", "Rheinisch-Bergischer-Kreis"), ("Stand", "27.09.2026 14:45")]
    for i, (bez, wert) in enumerate(kopfpaare, start=1):
        ws.cell(i, 1, bez)
        ws.cell(i, 2, wert)
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=7)
    for c, titel in enumerate(kopf, start=1):
        ws.cell(kopfzeile, c, titel)
    zeile = kopfzeile + 1
    for n, (vorname, nachname, geb, rolle, glied) in enumerate(personen):
        werte = [vorname, nachname, geb, rolle, glied, "aktiv",
                 "ja", "ja" if n % 3 else None, "ja" if n % 2 else None, "nein"]
        for v in range(len(VORAUSSETZUNGEN)):
            werte.append("ja" if (n + v) % 3 == 0 else "nein")
        for c, w in enumerate(werte, start=1):
            if w is not None:
                ws.cell(zeile, c, w)
        zeile += 1
    # Leerzeilen am Ende, wie sie Excel-Exporte hinterlassen: "" bzw. Leerzeichen statt None,
    # damit max_row steigt und der Import sie wirklich als leer erkennen muss.
    for _ in range(leerzeilen):
        for c in range(1, len(kopf) + 1):
            ws.cell(zeile, c, "" if c % 2 else "  ")
        zeile += 1
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_helfer():
    assert normalisiert("  Max\n\tMuster  MANN ") == "max muster mann"
    assert normalisiert(None) == ""
    k1 = schluessel("Anna", "Beispiel", "2002-10-03")
    assert k1 == ("anna", "beispiel", "2002-10-03")
    assert passt(k1, schluessel(" anna ", "BEISPIEL", None))          # fehlendes Datum ist Platzhalter
    assert passt(k1, schluessel("Anna", "Beispiel", "2002-10-03"))
    assert not passt(k1, schluessel("Anna", "Beispiel", "2001-01-01"))
    assert not passt(k1, schluessel("Anne", "Beispiel", None))
    assert schluessel("a", "b", "") == ("a", "b", None)

    # Datumsformate
    assert datum_lesen(datetime(2002, 10, 3, 14, 5)) == "2002-10-03"
    assert datum_lesen(date(2002, 10, 3)) == "2002-10-03"
    assert datum_lesen("2002-10-03 00:00:00") == "2002-10-03"
    assert datum_lesen("2002-10-03") == "2002-10-03"
    assert datum_lesen("03.10.2002") == "2002-10-03"
    assert datum_lesen("3.10.2002 18:00 Uhr") == "2002-10-03"
    assert datum_lesen("03.10.02") == "2002-10-03"
    assert datum_lesen("03.10.75") == "1975-10-03"
    assert datum_lesen("31.02.2002") is None
    assert datum_lesen("Hallo") is None and datum_lesen(None) is None and datum_lesen("") is None
    assert datum_lesen(17) is None                                        # Zahl ohne Freigabe: kein Datum
    assert datum_lesen(17, excel_zahl=True) is None                       # ein Alter ist kein Datum
    assert datum_lesen(37532, excel_zahl=True) == "2002-10-03"            # Excel-Seriennummer
    assert datum_lesen(True, excel_zahl=True) is None

    # Erfüllt / nicht erfüllt
    for w in ("ja", "J", "x", "✓", "✔", "wahr", "TRUE", "yes", "1", "ok", "erfüllt", "erfuellt", "vorhanden",
              1, True, "03.10.2024", "2024-10-03 00:00:00", datetime(2024, 1, 1), date(2024, 1, 1)):
        assert ist_erfuellt(w) is True, w
    for w in ("nein", "n", "-", "–", "falsch", "false", "no", "0", "offen", "fehlt", "leer", "", "  ", None,
              0, False, "irgendwas", 2):
        assert ist_erfuellt(w) is False, w

    # Zeitraum
    assert zeitraum_lesen("04.12.2026 18:00 Uhr bis 13.12.2026 19:00 Uhr") == ("2026-12-04", "2026-12-13")
    assert zeitraum_lesen("2026-12-04 – 2026-12-13") == ("2026-12-04", "2026-12-13")
    assert zeitraum_lesen("am 04.12.2026") == ("2026-12-04", "2026-12-04")
    assert zeitraum_lesen(datetime(2026, 12, 4, 18)) == ("2026-12-04", "2026-12-04")
    assert zeitraum_lesen("irgendwann") == (None, None)
    assert zeitraum_lesen(None) == (None, None)
    assert "voraussetzung" in ZUORDNUNGEN and "ignorieren" in ZUORDNUNGEN


def test_analysieren_beispieldatei():
    e = analysieren(beispiel_xlsx(), dateiname="2026-0008 - Strömungsretter 1.xlsx")

    # Kopfzeile und Kopfbereich (verbundene Zellen B:G, Wert steht in B)
    assert e["kopfzeile"] == 7
    assert e["lehrgang"] == {"titel": "Strömungsretter 1 (SR1)", "nummer": "2026-0008",
                             "datum_von": "2026-12-04", "datum_bis": "2026-12-13", "ort": "Rheinisch-Bergischer-Kreis"}

    # Spaltenzuordnung
    zu = {s["titel"]: s["zuordnung"] for s in e["spalten"]}
    assert zu["Vorname"] == "vorname" and zu["Nachname"] == "name"
    assert zu["Alter"] == "geburtsdatum"                      # Datumswerte in „Alter“
    assert zu["Rolle"] == "rolle" and zu["Gliederung"] == "gliederung"
    assert zu["Status"] == "extra"                            # „aktiv“ ist nicht boolesch
    for t in VERWALTUNG:
        assert zu[t] == "voraussetzung", t
    assert sum(1 for s in e["spalten"] if s["zuordnung"] == "voraussetzung") == 4 + 15
    assert [s["index"] for s in e["spalten"]] == list(range(len(KOPF)))
    # Titel mit Zeilenumbruch → Leerzeichen, Reihenfolge wie in der Datei
    assert len(e["voraussetzungen"]) == 19
    assert e["voraussetzungen"][0] == {"spaltenindex": 6, "bezeichnung": "Bestätigung TN"}
    # … auch wenn „Heimatgliederung“ im Titel steht: ja/nein-Werte machen die Spalte zur Voraussetzung
    assert e["voraussetzungen"][5]["bezeichnung"] == ("Mitgliedschaft in der DLRG (wird automatisch geprüft "
                                                      "bzw. bei der Heimatgliederung angefragt)")
    assert all("\n" not in v["bezeichnung"] for v in e["voraussetzungen"])
    assert all("\n" not in s["titel"] for s in e["spalten"])
    # Beispiele: höchstens 3, als Text, keine Dubletten
    sp = {s["titel"]: s for s in e["spalten"]}
    assert sp["Vorname"]["beispiele"] == ["Anna", "Bernd", "Clara"]
    assert sp["Zahlung"]["beispiele"] == ["nein"]
    assert "2005-06-30" in sp["Alter"]["beispiele"]

    # Teilnehmende: 7 Zeilen − 1 Lehrgangsleitung − 1 Duplikat = 5
    tn = e["teilnehmer"]
    assert [t["name"] for t in tn] == ["Beispiel", "Probe", "Test", "Fiktiv", "Erfunden"]
    assert [t["zeile"] for t in tn] == [8, 9, 10, 12, 14]
    anna = tn[0]
    assert anna["vorname"] == "Anna" and anna["geburtsdatum"] == "2002-10-03"
    assert anna["gliederung"] == "Musterstadt" and anna["email"] == "" and anna["bemerkung"] == ""
    assert anna["extra"] == {"Status": "aktiv", "Rolle": "Teilnehmender"}
    assert tn[2]["geburtsdatum"] == "2005-06-30"                       # Zelle mit echtem datetime
    # „Vorname Nachname“ in der Nachname-Spalte ohne Vorname
    assert tn[4]["vorname"] == "Frank" and tn[4]["name"] == "Erfunden"
    # Voraussetzungen: Schlüssel = str(Spaltenindex), alle 19 Spalten je TN, Werte bool
    assert set(anna["voraussetzungen"]) == {str(i) for i in range(6, 25)}
    assert all(isinstance(v, bool) for t in tn for v in t["voraussetzungen"].values())
    assert anna["voraussetzungen"]["6"] is True                       # „ja“
    assert anna["voraussetzungen"]["7"] is False                      # leere Zelle (n=0 → None)
    assert anna["voraussetzungen"]["9"] is False                      # „nein“
    assert anna["voraussetzungen"]["10"] is True and anna["voraussetzungen"]["11"] is False
    assert tn[1]["voraussetzungen"]["7"] is True

    # Rolle ≠ Teilnehmer → Ausbilder
    assert e["ausbilder"] == [{"name": "Dirk Muster", "funktion": "Lehrgangsleitung", "ist_leitung": True, "zeile": 11}]

    # Warnungen: Duplikat und Leerzeilen, als deutsche Sätze
    w = e["warnungen"]
    assert any("Zeile 13" in s and "Zeile 8" in s and "übersprungen" in s for s in w), w
    assert any("2 leere Zeilen übersprungen" in s for s in w), w
    assert len(w) == 2, w


def test_zuordnung_ueberschreiben():
    # Schlüssel als int ODER als Text (JSON), Voraussetzung → ignorieren, Extra → Voraussetzung
    e = analysieren(beispiel_xlsx(), zuordnung={6: "ignorieren", "5": "voraussetzung", "9": "bemerkung"})
    sp = {s["index"]: s["zuordnung"] for s in e["spalten"]}
    assert sp[6] == "ignorieren" and sp[5] == "voraussetzung" and sp[9] == "bemerkung"
    idx = [v["spaltenindex"] for v in e["voraussetzungen"]]
    assert 6 not in idx and 5 in idx and 9 not in idx and len(idx) == 18
    anna = e["teilnehmer"][0]
    assert "6" not in anna["voraussetzungen"] and anna["voraussetzungen"]["5"] is False   # „aktiv“ ist kein Haken
    assert anna["bemerkung"] == "nein" and "Status" not in anna["extra"]

    # Namensspalte auf ignorieren → ohne Nachnamen entsteht kein Teilnehmender; statt eines Fehlers
    # kommt eine korrigierbare Vorschau mit einem Hinweis auf die Zuordnung (nicht je Zeile einer).
    e2 = analysieren(beispiel_xlsx(), zuordnung={"1": "ignorieren"})
    assert e2["teilnehmer"] == [] and any("Nachname" in w for w in e2["warnungen"]), e2["warnungen"]
    assert sum("hat keinen Nachnamen" in w for w in e2["warnungen"]) == 0

    # Unbrauchbare Vorgaben werden gemeldet, nicht zum Fehler
    e3 = analysieren(beispiel_xlsx(), zuordnung={"abc": "name", "2": "quatsch"})
    assert any("abc" in s for s in e3["warnungen"]) and any("quatsch" in s for s in e3["warnungen"])
    assert {s["index"]: s["zuordnung"] for s in e3["spalten"]}[2] == "geburtsdatum"

    # Zwei Spalten als „name“: die erste gewinnt, die zweite wird weitere Angabe
    e4 = analysieren(beispiel_xlsx(), zuordnung={4: "name"})
    sp4 = {s["index"]: s["zuordnung"] for s in e4["spalten"]}
    assert sp4[1] == "name" and sp4[4] == "extra"
    assert e4["teilnehmer"][0]["extra"]["Gliederung"] == "Musterstadt"
    assert any("zweites Mal" in s for s in e4["warnungen"])


def test_freie_liste_und_sonderfaelle():
    """Von Hand gebaute Liste: Kopfzeile in Zeile 1, andere Synonyme, Zahlen-Alter, Excel-Seriennummer,
    Zeile ohne Namen mitten in der Liste, bytes statt Dateiobjekt."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Geb.", "E-Mail", "OG", "Alter", "Bemerkung", "Funktion", "Erste Hilfe", "Telefon", None, "Leer"])
    ws.append(["Max Mustermann", 37532, "max@example.org", "Musterstadt", 23, "kommt später", "", "x", "0123", "verirrt", None])
    ws.append(["Erika Musterfrau", "15.01.99", "", "Beispielhausen", 27, "", "Teilnehmerin", "-", None, None, None])
    ws.append(["", "", "", "Beispielhausen", None, "ohne Namen", "", "", None, None, None])
    ws.append(["Referentin Extern", date(1980, 4, 1), "", "", 46, "", "Referent", "ja", None, None, None])
    ws.append(["", " ", "", None, None, None, None, None, None, None, None])   # Leerzeile mit "" – None erzeugte keine Zelle
    buf = io.BytesIO()
    wb.save(buf)
    e = analysieren(buf.getvalue())                         # bytes ohne Dateiname
    assert e["kopfzeile"] == 1
    assert e["lehrgang"] == {"titel": None, "nummer": None, "datum_von": None, "datum_bis": None, "ort": None}
    zu = {s["titel"]: s["zuordnung"] for s in e["spalten"]}
    assert zu == {"Name": "name", "Geb.": "geburtsdatum", "E-Mail": "email", "OG": "gliederung", "Alter": "extra",
                  "Bemerkung": "bemerkung", "Funktion": "rolle", "Erste Hilfe": "voraussetzung", "Telefon": "extra",
                  "": "ignorieren", "Leer": "extra"}
    # Spalte ohne Titel mit Daten wird gelistet (ignorieren), Spalte „Leer“ ohne Werte bleibt extra
    ohne_titel = [s for s in e["spalten"] if s["titel"] == ""]
    assert len(ohne_titel) == 1 and ohne_titel[0]["index"] == 9 and ohne_titel[0]["beispiele"] == ["verirrt"]
    tn = e["teilnehmer"]
    assert len(tn) == 2
    assert tn[0]["vorname"] == "Max" and tn[0]["name"] == "Mustermann" and tn[0]["geburtsdatum"] == "2002-10-03"
    assert tn[0]["email"] == "max@example.org" and tn[0]["gliederung"] == "Musterstadt"
    assert tn[0]["extra"] == {"Alter": "23", "Telefon": "0123"}
    assert tn[0]["bemerkung"] == "kommt später"
    assert tn[0]["voraussetzungen"] == {"7": True}
    assert tn[1]["vorname"] == "Erika" and tn[1]["geburtsdatum"] == "1999-01-15"
    assert tn[1]["voraussetzungen"] == {"7": False}
    assert tn[1]["extra"]["Funktion"] == "Teilnehmerin"        # Rollentext bleibt als Zusatzangabe
    assert e["ausbilder"] == [{"name": "Referentin Extern", "funktion": "Referent", "ist_leitung": False, "zeile": 5}]
    w = e["warnungen"]
    assert any("Zeile 4" in s and "Nachnamen" in s and "übersprungen" in s for s in w), w
    assert any("Eine leere Zeile" in s for s in w), w

    # Synonym mitten im Titel zählt nur bei nicht-booleschen Werten: „DLRG-Gliederung“ mit Ortsnamen
    # → gliederung, „Bestätigung Gliederung“ mit ja/nein → Voraussetzung
    wb5 = openpyxl.Workbook()
    ws5 = wb5.active
    ws5.append(["Vorname", "Nachname", "DLRG-Gliederung", "Bestätigung Gliederung", "Kontakt-Mail"])
    ws5.append(["Max", "Mustermann", "Musterstadt", "ja", "max@example.org"])
    ws5.append(["Erika", "Musterfrau", "Beispielhausen", "nein", ""])
    b5 = io.BytesIO()
    wb5.save(b5)
    e5 = analysieren(b5)
    assert [s["zuordnung"] for s in e5["spalten"]] == ["vorname", "name", "gliederung", "voraussetzung", "email"]
    assert e5["teilnehmer"][0]["gliederung"] == "Musterstadt" and e5["teilnehmer"][0]["email"] == "max@example.org"

    # Unlesbares Geburtsdatum → Warnung, Feld bleibt leer
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.append(["Vorname", "Nachname", "Geburtsdatum"])
    ws2.append(["Ute", "Unklar", "irgendwann"])
    b2 = io.BytesIO()
    wb2.save(b2)
    e2 = analysieren(b2)
    assert e2["teilnehmer"][0]["geburtsdatum"] is None
    assert any("Zeile 2" in s and "irgendwann" in s for s in e2["warnungen"])

    # Kopfzeile ohne Namensspalte: die Zeile mit den meisten Titeln wird genommen, mit Warnung
    wb3 = openpyxl.Workbook()
    ws3 = wb3.active
    ws3.append(["Liste", None])
    ws3.append(["Teilnehmende", "Verein", "Erste Hilfe"])
    ws3.append(["Max Mustermann", "Musterstadt", "ja"])
    b3 = io.BytesIO()
    wb3.save(b3)
    e3 = analysieren(b3, zuordnung={0: "name"})
    assert e3["kopfzeile"] == 2 and any("Kopfzeile" in s for s in e3["warnungen"])
    assert e3["teilnehmer"][0]["name"] == "Mustermann" and e3["teilnehmer"][0]["gliederung"] == "Musterstadt"

    # Erstes Arbeitsblatt leer → das erste mit Daten zählt
    wb4 = openpyxl.Workbook()
    wb4.active.title = "Deckblatt"
    ws4 = wb4.create_sheet("Liste")
    ws4.append(["Vorname", "Nachname", "Gliederung"])
    ws4.append(["Max", "Mustermann", "Musterstadt"])
    b4 = io.BytesIO()
    wb4.save(b4)
    assert analysieren(b4)["teilnehmer"][0]["name"] == "Mustermann"


def test_importfehler():
    with pytest.raises(ImportFehler):
        analysieren(beispiel_xlsx(), dateiname="liste.xls")             # falsche Endung
    with pytest.raises(ImportFehler):
        analysieren(beispiel_xlsx(), dateiname="liste.csv")
    with pytest.raises(ImportFehler):
        analysieren(b"Vorname;Nachname\nMax;Mustermann\n", dateiname="liste.xlsx")   # kein ZIP/xlsx
    with pytest.raises(ImportFehler):
        analysieren(b"")
    with pytest.raises(ImportFehler):
        analysieren(b"PK\x03\x04kaputt", dateiname="liste.xlsx")         # beschädigt
    # Keine Kopfzeile
    wb = openpyxl.Workbook()
    wb.active.append(["nur", "zwei"])
    b = io.BytesIO()
    wb.save(b)
    with pytest.raises(ImportFehler):
        analysieren(b)
    # Ganz leeres Blatt
    wb = openpyxl.Workbook()
    b = io.BytesIO()
    wb.save(b)
    with pytest.raises(ImportFehler):
        analysieren(b)
    # Kopfzeile, aber keine Datenzeile
    with pytest.raises(ImportFehler):
        analysieren(beispiel_xlsx(personen=[], leerzeilen=1))
    # Die Meldungen sind deutsche Sätze für die Oberfläche
    try:
        analysieren(beispiel_xlsx(), dateiname="liste.xls")
    except ImportFehler as e:
        assert ".xlsx" in str(e)


def test_leitungsrolle():
    """Leitung ist ein eigenes Wort in der Rolle, kein Teilstring: „Begleitung“ leitet nicht."""
    for rolle, erwartet in (("Lehrgangsleitung", True), ("Leitung", True), ("Leiter:in", True), ("Leiterin", True),
                            ("Leiter/innen", True), ("Leiter*in", True), ("Lehrgangsleiter", True),
                            ("Leitung und Prüfer", True), ("Leitung/Prüfer", True), ("Ltg.", True), ("LGL", True),
                            ("Co-Leitung", True), ("stellv. Leitung", True),
                            ("Begleitung", False), ("Anleitung Knoten", False), ("Ausbildungsleiter", False),
                            ("Schwimmleiter", False), ("Referent", False), ("Teilnehmender", False), ("", False),
                            (None, False)):
        assert ist_leitungsrolle(rolle) is erwartet, rolle
    # In der Datei: „Begleitung“ wird über das Ausbilderwort „leitung“ zwar Ausbilderin, aber keine Leitung.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Vorname", "Nachname", "Rolle"])
    ws.append(["Ida", "Import", "TN"])
    ws.append(["Berta", "Begleit", "Begleitung"])
    ws.append(["Lena", "Leit", "Leitung und Prüferin"])
    b = io.BytesIO()
    wb.save(b)
    e = analysieren(b)
    assert [(a["name"], a["ist_leitung"]) for a in e["ausbilder"]] == [("Berta Begleit", False), ("Lena Leit", True)]
    assert [t["name"] for t in e["teilnehmer"]] == ["Import"]


def test_grenzen():
    """Dateien, die klein aussehen und beim Lesen explodieren, werden mit Klartext abgewiesen – geprüft
    mit kleinen Dateien, die eine Grenze reißen, ohne den Test zu bremsen."""
    def mit_zelle(zeile, spalte):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Vorname", "Nachname"])
        ws.append(["Max", "Muster"])
        ws.cell(zeile, spalte, "x")
        b = io.BytesIO()
        wb.save(b)
        return b.getvalue()

    # Eine einzige Zelle weit außerhalb: höchste Zeile, höchste Spalte, Rechteck. openpyxl legte für
    # jedes Feld des Rechtecks ein Zellobjekt an; die Grenzen gelten der Form, nicht der Dateigröße –
    # jede dieser Dateien ist unter 10 KB.
    for zeile, spalte, hinweis in ((30_000, 1, "Zeile 30000"), (3, 400, "Spalte 400"), (5_000, 200, "Felder")):
        roh = mit_zelle(zeile, spalte)
        assert len(roh) < 10_000
        with pytest.raises(ImportFehler) as e:
            analysieren(roh)
        assert hinweis in str(e.value), (zeile, spalte, str(e.value))
    assert analysieren(mit_zelle(1_000, 50))["teilnehmer"][0]["name"] == "Muster"       # eine normale Liste

    # Zu viele Arbeitsblätter – alle würden beim Öffnen geladen.
    wb = openpyxl.Workbook()
    for i in range(pi.MAX_ARBEITSBLAETTER):                 # plus das vorhandene: eins zu viel
        wb.create_sheet(f"Blatt {i}")
    wb.worksheets[-1].append(["Vorname", "Nachname"])
    wb.worksheets[-1].append(["Max", "Muster"])
    b = io.BytesIO()
    wb.save(b)
    with pytest.raises(ImportFehler) as e:
        analysieren(b)
    assert "Arbeitsblätter" in str(e.value)

    # Das Archiv selbst: zu viele Einträge; zu groß entpackt (die Grenze dafür künstlich klein, damit
    # der Test keine 200 MB schreibt – geprüft wird der Weg, nicht die Zahl).
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as zf:
        for i in range(pi.MAX_ZIP_EINTRAEGE + 1):
            zf.writestr(f"e{i}", b"x")
    with pytest.raises(ImportFehler) as e:
        analysieren(b.getvalue())
    assert "Einträgen" in str(e.value)
    alt = pi.MAX_ENTPACKT
    pi.MAX_ENTPACKT = 1024
    try:
        with pytest.raises(ImportFehler) as e:
            analysieren(beispiel_xlsx())
        assert "entpackt" in str(e.value)
    finally:
        pi.MAX_ENTPACKT = alt
    assert analysieren(beispiel_xlsx())["kopfzeile"] == 7                              # danach wieder lesbar

    # Mehr Datenzeilen als auf einmal übernommen werden: Fehler mit Zahl; Leerzeilen zählen nicht.
    def liste(n, leerzeilen=0):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Vorname", "Nachname"])
        for i in range(n):
            ws.append([f"V{i}", f"N{i}"])
        for _ in range(leerzeilen):
            ws.append(["", "  "])
        b = io.BytesIO()
        wb.save(b)
        return b
    with pytest.raises(ImportFehler) as e:
        analysieren(liste(pi.MAX_DATENZEILEN + 1))
    assert f"{pi.MAX_DATENZEILEN + 1} Datenzeilen" in str(e.value)
    assert len(analysieren(liste(pi.MAX_DATENZEILEN, leerzeilen=5))["teilnehmer"]) == pi.MAX_DATENZEILEN

    # Weitere Angaben: höchstens 50 Spalten (der Rest wird „ignorieren“ und einmal gemeldet), Titel auf
    # 100 und Werte auf 500 Zeichen gekürzt.
    kopf = ["Vorname", "Nachname", "T" * 150] + [f"Z{i}" for i in range(pi.MAX_EXTRA_SPALTEN + 4)]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(kopf)
    ws.append(["Max", "Muster", "w" * 600] + [f"wert{i}" for i in range(pi.MAX_EXTRA_SPALTEN + 4)])
    b = io.BytesIO()
    wb.save(b)
    e = analysieren(b)
    zu = [s["zuordnung"] for s in e["spalten"]]
    assert zu.count("extra") == pi.MAX_EXTRA_SPALTEN and zu[-5:] == ["ignorieren"] * 5, zu
    assert sum("weiteren Angaben" in w for w in e["warnungen"]) == 1, e["warnungen"]
    extra = e["teilnehmer"][0]["extra"]
    assert len(extra) == pi.MAX_EXTRA_SPALTEN and extra["T" * pi.MAX_EXTRA_TITEL] == "w" * pi.MAX_EXTRA_WERT


def xlsx_bomben():
    """Kleine Arbeitsmappen (je ein paar KB), die openpyxl beim LADEN aufbläht – Sicherheitsprüfung zum Bestandsimport
    (Material, Phase 2 · L3, Befund 1); gemessen vorher 1,6–2,8 GB und bis 76 s je Datei, „riesig“ unbegrenzt:
    ein Bereich verbundener Zellen, ein Link und ein Kommentar über A3:KN20000 (6 Mio. Felder innerhalb der Grenzen für
    Zeilen und Spalten), verbundene Zellen bis XFD1048576 und 9,5 MB dichte leere Zellen „<c/>“. Dazu Verkleidungen,
    an denen eine Textsuche vorbeiliefe: Namensraum-Präfix, Zeichenreferenz im Bereich, UTF-16, eine Entität, die
    leere Zellen vervielfacht, ein Blatt unter fremdem Pfad und ein Link über ganze Spalten. Kopfzeile mit Vorname,
    Nachname, Hersteller, Modell und Seriennummer – beide Importe (Prüfungen, Bestandsimport) erkennen sie."""
    from openpyxl.comments import Comment

    def grund(kommentar=False):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Vorname", "Nachname", "Hersteller", "Modell", "Seriennummer"])
        ws.append(["Max", "Muster", "Testwerk", "HMS Beta", "SN-0001"])
        if kommentar:
            ws["A1"].comment = Comment("Platzhalter", "Platzhalter")
        b = io.BytesIO()
        wb.save(b)
        return b.getvalue()

    def umbauen(roh, aendern, teil="xl/worksheets/sheet1.xml", umbenennen=None):
        zin, out = zipfile.ZipFile(io.BytesIO(roh)), io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
            for e in zin.infolist():
                name, daten = e.filename, zin.read(e.filename)
                if name == teil:
                    daten = aendern(daten.decode())
                    daten = daten if isinstance(daten, bytes) else daten.encode()
                if umbenennen:
                    name, daten = umbenennen(name, daten)
                zo.writestr(name, daten)
        return out.getvalue()

    def danach(teil):
        return lambda x: x.replace("</sheetData>", "</sheetData>" + teil)
    merge = '<mergeCells count="1"><mergeCell ref="A3:KN20000"/></mergeCells>'
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    return {
        "merge_6m": umbauen(grund(), danach(merge)),
        "link_6m": umbauen(grund(), danach('<hyperlinks><hyperlink ref="A3:KN20000" display="x" location="A1"/>'
                                           "</hyperlinks>")),
        "kommentar_6m": umbauen(grund(True), lambda x: x.replace('ref="A1"', 'ref="A3:KN20000"'),
                                "xl/comments/comment1.xml"),
        "merge_riesig": umbauen(grund(), danach('<mergeCells count="1"><mergeCell ref="A3:XFD1048576"/></mergeCells>')),
        "leere_c": umbauen(grund(), lambda x: x.replace("</sheetData>", ("<row>" + "<c/>" * 290 + "</row>") * 8121
                                                        + "</sheetData>")),
        "praefix": umbauen(grund(), danach(f'<x:mergeCells xmlns:x="{ns}" count="1"><x:mergeCell ref="A3:KN20000"/>'
                                           "</x:mergeCells>")),
        "zeichenreferenz": umbauen(grund(), danach(merge.replace("A3:KN", "A3&#58;KN"))),
        "utf16": umbauen(grund(), lambda x: danach(merge)(x).encode("utf-16")),
        "entitaet": umbauen(grund(), lambda x: x.replace(
            "<worksheet", '<!DOCTYPE worksheet [<!ENTITY z "' + "<c/>" * 1000 + '"><!ENTITY y "' + "&z;" * 1000
            + '">]><worksheet', 1).replace("</sheetData>", "<row>" + "&y;" * 6 + "</row></sheetData>")),
        "fremder_pfad": umbauen(grund(), danach(merge), umbenennen=lambda n, d: (
            ("xl/daten/blatt.dat", d) if n == "xl/worksheets/sheet1.xml" else
            (n, d.replace(b"worksheets/sheet1.xml", b"daten/blatt.dat")) if n == "xl/_rels/workbook.xml.rels"
            else (n, d))),
        "ganze_spalten": umbauen(grund(), lambda x: danach('<hyperlinks><hyperlink ref="A:KN" display="x" '
                                                           'location="A1"/></hyperlinks>')(x.replace(
                                                               "</sheetData>", '<row r="20000"><c r="A20000" '
                                                               't="inlineStr"><is><t>x</t></is></c></row></sheetData>'))),
    }


def test_grenzen_vor_dem_laden():
    """Was openpyxl schon beim Laden anlegt (Zellen je <c>, Felder verbundener Zellen, Links, Kommentare), wird vorher
    gezählt (_blaetter_vorpruefen): Jede Datei aus xlsx_bomben wird abgelehnt, ohne dass openpyxl lädt – in unter
    2 s und mit wenig Speicher. Zulässige Listen mit verbundenen Kopfzellen und Links bleiben lesbar."""
    geladen = []
    alt = pi.openpyxl.load_workbook
    pi.openpyxl.load_workbook = lambda *a, **k: geladen.append(1) or alt(*a, **k)
    try:
        for name, roh in xlsx_bomben().items():
            assert len(roh) < 40_000, (name, len(roh))
            tracemalloc.start()
            t0 = time.time()
            with pytest.raises(ImportFehler) as e:
                analysieren(roh, dateiname=f"{name}.xlsx")
            dauer, spitze = time.time() - t0, tracemalloc.get_traced_memory()[1]
            tracemalloc.stop()
            text = str(e.value)
            assert not geladen and dauer < 2 and spitze < 50 * 1024 * 1024, (name, dauer, spitze, text)
            soll = ("bis Zeile 1048576 und Spalte 16384" if name == "merge_riesig" else
                    "konnte nicht gelesen werden" if name == "ganze_spalten" else "mehr als 500.000 Felder")
            assert soll in text, (name, text)
        # Zu viele Elemente insgesamt (Grenze künstlich klein – geprüft wird der Weg, nicht die Zahl).
        alt_max = pi.MAX_XML_ELEMENTE
        pi.MAX_XML_ELEMENTE = 100
        try:
            with pytest.raises(ImportFehler) as e:
                analysieren(beispiel_xlsx())
            assert "zu umfangreich" in str(e.value) and not geladen
        finally:
            pi.MAX_XML_ELEMENTE = alt_max
        # Zulässig: Kopfpaare mit verbundenen Zellen (beispiel_xlsx) und eine Liste mit Links.
        assert analysieren(beispiel_xlsx())["kopfzeile"] == 7 and geladen
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Vorname", "Nachname", "Link"])
        for n in range(20):
            ws.append([f"V{n}", f"N{n}", "Wiki"])
            ws.cell(n + 2, 3).hyperlink = "https://example.org/platzhalter"
        ws.merge_cells("D1:G1")
        b = io.BytesIO()
        wb.save(b)
        assert len(analysieren(b)["teilnehmer"]) == 20
    finally:
        pi.openpyxl.load_workbook = alt


if __name__ == "__main__":
    test_helfer()
    test_analysieren_beispieldatei()
    test_zuordnung_ueberschreiben()
    test_freie_liste_und_sonderfaelle()
    test_importfehler()
    test_leitungsrolle()
    test_grenzen()
    test_grenzen_vor_dem_laden()
    print("Import-Test bestanden.")
